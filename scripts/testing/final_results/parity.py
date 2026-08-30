"""Semantic parity comparison for canonical Markdown and generated DOCX reports."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document


_HEADING = re.compile(r"^#{1,6}\s+(.+?)\s*$")
_SEPARATOR = re.compile(r"^:?-{3,}:?$")


def _split_markdown_row(line: str) -> list[str]:
    content = line.strip()[1:-1]
    cells: list[str] = []
    cell: list[str] = []
    index = 0
    while index < len(content):
        character = content[index]
        if character == "\\" and index + 1 < len(content) and content[index + 1] == "|":
            cell.append("|")
            index += 2
            continue
        if character == "|":
            cells.append("".join(cell).strip().replace("<br>", "\n"))
            cell = []
        else:
            cell.append(character)
        index += 1
    cells.append("".join(cell).strip().replace("<br>", "\n"))
    return cells


def _markdown_semantics(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    headings: list[str] = []
    tables: list[list[list[str]]] = []
    index = 0
    while index < len(lines):
        heading = _HEADING.match(lines[index].strip())
        if heading:
            headings.append(heading.group(1))
            index += 1
            continue
        if lines[index].strip().startswith("|") and lines[index].strip().endswith("|"):
            table: list[list[str]] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                row = _split_markdown_row(candidate)
                if not all(_SEPARATOR.fullmatch(cell) for cell in row):
                    table.append(row)
                index += 1
            tables.append(table)
            continue
        index += 1
    return headings, tables


def _docx_semantics(path: Path) -> tuple[list[str], list[list[list[str]]]]:
    document = Document(path)
    headings = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style is not None
        and (paragraph.style.name == "Title" or paragraph.style.name.startswith("Heading "))
    ]
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]
    return headings, tables


def compare_markdown_docx(markdown: Path, docx: Path) -> dict[str, object]:
    """Compare heading order/text and every table cell across canonical formats."""
    markdown_headings, markdown_tables = _markdown_semantics(Path(markdown))
    docx_headings, docx_tables = _docx_semantics(Path(docx))
    headings_match = markdown_headings == docx_headings
    tables_match = markdown_tables == docx_tables
    differences: list[dict[str, object]] = []
    if not headings_match:
        differences.append(
            {"kind": "headings", "markdown": markdown_headings, "docx": docx_headings}
        )
    if not tables_match:
        differences.append(
            {"kind": "tables", "markdown": markdown_tables, "docx": docx_tables}
        )
    return {
        "matches": headings_match and tables_match,
        "headings_match": headings_match,
        "tables_match": tables_match,
        "headings": {"markdown": markdown_headings, "docx": docx_headings},
        "tables": {"markdown": markdown_tables, "docx": docx_tables},
        "differences": differences,
    }
