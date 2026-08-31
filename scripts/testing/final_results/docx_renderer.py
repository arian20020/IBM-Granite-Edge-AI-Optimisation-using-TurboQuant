"""Professional, deterministic DOCX rendering for canonical final-results reports."""

from __future__ import annotations

import tempfile
import zipfile
from datetime import datetime, time
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Inches, Pt, RGBColor
from docx.table import _Cell, _Row, Table
from docx.text.paragraph import Paragraph

from .report_model import Report, ReportNote, ReportParagraph, ReportTable


NAVY = "0F2747"
TEAL = "0F766E"
BLUE = "0B63CE"
WHITE = RGBColor(255, 255, 255)
PORTRAIT_WIDTH = Inches(8.27)
PORTRAIT_HEIGHT = Inches(11.69)

_STATUS_FILLS = {
    "passed": "E2F0D9",
    "failed": "FCE4D6",
    "blocked": "FFF2CC",
    "not collected": "E7E6E6",
    "unavailable": "E7E6E6",
    "artifact unavailable": "E7E6E6",
    "not applicable": "E7E6E6",
    "verified": "DDEBF7",
}


def _set_fill(properties, colour: str) -> None:
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), colour)


def _shade_paragraph(paragraph: Paragraph, colour: str) -> None:
    _set_fill(paragraph._p.get_or_add_pPr(), colour)


def _shade_cell(cell: _Cell, colour: str) -> None:
    _set_fill(cell._tc.get_or_add_tcPr(), colour)


def _prevent_row_split(row: _Row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _repeat_header(row: _Row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        properties.append(header)


def _set_portrait(section: Section) -> None:
    section.page_width = PORTRAIT_WIDTH
    section.page_height = PORTRAIT_HEIGHT
    page_size = section._sectPr.get_or_add_pgSz()
    page_size.attrib.pop(qn("w:orient"), None)


def _set_landscape(section: Section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = PORTRAIT_HEIGHT
    section.page_height = PORTRAIT_WIDTH


def _configure_section(section: Section, *, landscape: bool) -> None:
    if landscape:
        _set_landscape(section)
    else:
        _set_portrait(section)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)


def _add_page_field(paragraph: Paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _configure_footer(section: Section, report: Report) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.8), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run(f"{report.route_id} | Revision {report.revision}")
    paragraph.add_run("\tPage ")
    _add_page_field(paragraph)
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(89, 89, 89)


def _configure_styles(document: DocumentObject) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    for name, size, colour in (
        ("Title", 22, WHITE),
        ("Heading 1", 15, WHITE),
        ("Heading 2", 11, WHITE),
    ):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = colour
    document.styles["Subtitle"].font.name = "Aptos"
    document.styles["Subtitle"].font.size = Pt(9)


def _add_band_heading(document: DocumentObject, text: str, style: str, colour: str) -> Paragraph:
    paragraph = document.add_paragraph(text, style=style)
    _shade_paragraph(paragraph, colour)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def _validate_table(table: ReportTable) -> None:
    width = len(table.columns)
    for row in table.rows:
        if len(row) != width:
            raise ValueError(
                f"table {table.table_id!r} has {width} columns but a row has {len(row)} cells"
            )


def _estimated_table_characters(table: ReportTable) -> int:
    if not table.columns:
        return 0
    return sum(
        max([len(table.columns[index]), *(len(row[index]) for row in table.rows)])
        for index in range(len(table.columns))
    )


def _is_wide(table: ReportTable) -> bool:
    return len(table.columns) > 6 or _estimated_table_characters(table) > 95


def _write_cell(cell: _Cell, value: str) -> None:
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)


def _style_table(table: Table) -> None:
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        _prevent_row_split(row)
        if row_index == 0:
            _repeat_header(row)
        for cell in row.cells:
            if row_index == 0:
                _shade_cell(cell, BLUE)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = WHITE
            elif row_index % 2 == 0:
                _shade_cell(cell, "EEF4F8")
            status_fill = _STATUS_FILLS.get(cell.text.strip().casefold())
            if status_fill is not None:
                _shade_cell(cell, status_fill)


def _add_table(document: DocumentObject, columns: tuple[str, ...], rows: tuple[tuple[str, ...], ...]) -> None:
    word_table = document.add_table(rows=1 + len(rows), cols=len(columns))
    for column_index, value in enumerate(columns):
        _write_cell(word_table.cell(0, column_index), value)
    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            _write_cell(word_table.cell(row_index, column_index), value)
    _style_table(word_table)


def _add_report_table(document: DocumentObject, table: ReportTable) -> None:
    _add_band_heading(document, f"{table.table_id} — {table.title}", "Heading 2", TEAL)
    if table.subtitle:
        subtitle = document.add_paragraph(table.subtitle, style="Subtitle")
        subtitle.paragraph_format.keep_with_next = True
    _add_table(document, table.columns, table.rows)
    for footnote in table.footnotes:
        paragraph = document.add_paragraph(footnote, style="Caption")
        paragraph.paragraph_format.space_after = Pt(3)


def _add_document_control(document: DocumentObject, report: Report) -> None:
    rows = (
        ("Route ID", report.route_id),
        ("Revision", report.revision),
        ("Generated date", report.generated_date.isoformat()),
        ("Evidence IDs", ", ".join(report.evidence_ids)),
    )
    _add_table(document, ("Document control", "Value"), rows)


def _normalize_package(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        import os

        os.close(descriptor)
        with zipfile.ZipFile(
            temporary_name, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for original, data in sorted(members, key=lambda member: member[0].filename):
                info = zipfile.ZipInfo(original.filename, (1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = original.external_attr
                info.internal_attr = original.internal_attr
                info.create_system = original.create_system
                info.comment = original.comment
                target.writestr(info, data)
        Path(temporary_name).replace(path)
    finally:
        Path(temporary_name).unlink(missing_ok=True)


def render_docx(report: Report, output: Path) -> None:
    """Render ``report`` to a deterministic, professionally styled Word document."""
    output = Path(output)
    document = Document()
    fixed_timestamp = datetime.combine(report.generated_date, time.min)
    document.core_properties.author = "IBM Granite Edge AI testing"
    document.core_properties.last_modified_by = "IBM Granite Edge AI testing"
    document.core_properties.created = fixed_timestamp
    document.core_properties.modified = fixed_timestamp
    document.core_properties.revision = 1
    _configure_styles(document)
    _configure_section(document.sections[0], landscape=False)
    _configure_footer(document.sections[0], report)

    title = document.add_paragraph(report.title, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _shade_paragraph(title, NAVY)
    title.paragraph_format.space_after = Pt(10)
    _add_document_control(document, report)

    landscape = False
    for section in report.sections:
        if landscape:
            word_section = document.add_section(WD_SECTION.NEW_PAGE)
            _configure_section(word_section, landscape=False)
            landscape = False
        _add_band_heading(document, section.title, "Heading 1", NAVY)
        for block in section.blocks:
            if isinstance(block, ReportTable):
                _validate_table(block)
                requested_landscape = _is_wide(block)
                if requested_landscape != landscape:
                    word_section = document.add_section(WD_SECTION.NEW_PAGE)
                    _configure_section(word_section, landscape=requested_landscape)
                    landscape = requested_landscape
                _add_report_table(document, block)
            else:
                if landscape:
                    word_section = document.add_section(WD_SECTION.NEW_PAGE)
                    _configure_section(word_section, landscape=False)
                    landscape = False
                if isinstance(block, ReportParagraph):
                    document.add_paragraph(block.text)
                elif isinstance(block, ReportNote):
                    note = document.add_paragraph()
                    label = note.add_run("Note: ")
                    label.bold = True
                    note.add_run(block.text)
                    _shade_paragraph(note, "E8F3F2")
                else:  # pragma: no cover - defensive guard for malformed callers.
                    raise TypeError(f"unsupported report block: {type(block)!r}")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    _normalize_package(output)
