"""Validate the Workbook 05 post-C workbook pack as untrusted data.

The validator is intentionally repository-only. It validates Markdown, CSV,
JSON and DOCX structure and cross-record consistency. It never executes a
captured command, accesses a model, or contacts a self-hosted runner.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Iterable, Mapping, Sequence

from docx import Document


@dataclass(frozen=True)
class PostCWorkbook:
    """One controlled workbook and the minimum structure it must retain."""

    workbook_id: str
    stage_id: str
    template_path: str
    generated_path: str
    required_headings: tuple[str, ...]
    minimum_table_count: int


@dataclass(frozen=True)
class PackIssue:
    """One stable validation failure suitable for CI and review reports."""

    code: str
    path: str
    message: str


POST_C_WORKBOOKS: tuple[PostCWorkbook, ...] = (
    PostCWorkbook(
        "WB-07",
        "D1",
        "docs/testing/workbooks/text-templates/07_Workbook_05_D1_Codec_Conformance_Controlled_Workbook_v1.md",
        "docs/testing/workbooks/generated/07_Workbook_05_D1_Codec_Conformance_Controlled_Workbook_v1.docx",
        ("Purpose", "Control boundary", "Initial revision record"),
        5,
    ),
    PostCWorkbook(
        "WB-08",
        "D2",
        "docs/testing/workbooks/text-templates/08_Workbook_05_D2_Diagnostic_KV_Sweep_Controlled_Workbook_v1.md",
        "docs/testing/workbooks/generated/08_Workbook_05_D2_Diagnostic_KV_Sweep_Controlled_Workbook_v1.docx",
        ("Purpose", "Control boundary", "Initial revision record"),
        5,
    ),
    PostCWorkbook(
        "WB-09",
        "E1",
        "docs/testing/workbooks/text-templates/09_Workbook_05_E1_Granite_3B_Frontier_Formal_Controlled_Workbook_v1.md",
        "docs/testing/workbooks/generated/09_Workbook_05_E1_Granite_3B_Frontier_Formal_Controlled_Workbook_v1.docx",
        ("Purpose", "Control boundary", "Initial revision record"),
        5,
    ),
    PostCWorkbook(
        "WB-10",
        "E2",
        "docs/testing/workbooks/text-templates/10_Workbook_05_E2_Granite_8B_Feasibility_Controlled_Workbook_v1.md",
        "docs/testing/workbooks/generated/10_Workbook_05_E2_Granite_8B_Feasibility_Controlled_Workbook_v1.docx",
        ("Purpose", "Control boundary", "Initial revision record"),
        5,
    ),
    PostCWorkbook(
        "WB-11",
        "E3",
        "docs/testing/workbooks/text-templates/11_Workbook_05_E3_Granite_30B_Bounded_Feasibility_Controlled_Workbook_v1.md",
        "docs/testing/workbooks/generated/11_Workbook_05_E3_Granite_30B_Bounded_Feasibility_Controlled_Workbook_v1.docx",
        ("Purpose", "Control boundary", "Initial revision record"),
        5,
    ),
    PostCWorkbook(
        "WB-12",
        "E4",
        "docs/testing/workbooks/text-templates/12_Workbook_05_E4_Cross_Family_Repeatability_Controlled_Workbook_v1.md",
        "docs/testing/workbooks/generated/12_Workbook_05_E4_Cross_Family_Repeatability_Controlled_Workbook_v1.docx",
        ("Purpose", "Control boundary", "Initial revision record"),
        5,
    ),
    PostCWorkbook(
        "WB-13",
        "F",
        "docs/testing/workbooks/text-templates/13_Workbook_05_F_Evidence_Closure_Controlled_Workbook_v1.md",
        "docs/testing/workbooks/generated/13_Workbook_05_F_Evidence_Closure_Controlled_Workbook_v1.docx",
        ("Purpose", "Control boundary", "Initial revision record"),
        5,
    ),
)

PACK_MANIFEST_PATH = Path(
    "docs/testing/workbooks/Workbook-05-Post-C-Pack-Manifest-v1.json"
)
CONTROLLED_MANIFEST_PATH = Path(
    "docs/testing/workbooks/Controlled-Workbook-Manifest.csv"
)
EXECUTION_INDEX_PATH = Path(
    "docs/testing/Workbook-05-Post-C-Execution-Index-v1.csv"
)
CONFIGURATION_MATRIX_PATH = Path(
    "docs/testing/Workbook-05-Post-C-Configuration-Matrix-v1.csv"
)
EVIDENCE_REGISTER_PATH = Path(
    "docs/testing/Workbook-05-Post-C-Evidence-Register-v1.csv"
)
REVISION_REGISTER_PATH = Path(
    "docs/testing/Workbook-05-Post-C-Revision-Register-v1.csv"
)

EXPECTED_STAGES = {"D1", "D2", "E1", "E2", "E3", "E4", "F"}
EXPECTED_STAGE_DEPENDENCIES: Mapping[str, list[str]] = {
    "D1": ["C1", "C2", "C3", "C4", "C5"],
    "D2": ["D1"],
    "E1": ["D2"],
    "E2": ["E1"],
    "E3": ["D2", "E1"],
    "E4": ["D2", "E1"],
    "F": ["D1", "D2", "E1", "E2", "E3", "E4"],
}
EXPECTED_STATUSES = {
    "Not started",
    "Blocked pending predecessor",
    "Passed",
    "Failed",
    "Blocked",
    "Not applicable",
    "Skipped by frontier",
    "Infrastructure interrupted",
}
EXPECTED_INDEX_HEADERS = {
    "Execution_Record_ID",
    "Stage_ID",
    "Test_ID",
    "Workbook_ID",
    "Route_ID",
    "Category",
    "Test_Title",
    "Model_ID",
    "Model_Scale",
    "Weight_Format",
    "Weight_Precision",
    "K_Algorithm",
    "K_Precision",
    "V_Algorithm",
    "V_Precision",
    "Context_Tokens",
    "Output_Tokens",
    "Repetition_Role",
    "Execution_Order",
    "Predecessor_Gate",
    "Expected_K_Record_Bytes",
    "Expected_V_Record_Bytes",
    "Verified_K_Record_Bytes",
    "Verified_V_Record_Bytes",
    "Memory_Rank",
    "Frontier_Status",
    "Formal_Statistics_Allowed",
    "Discovery_Only",
    "Status",
    "Blocker_Code",
    "Run_ID",
    "Artifact_ID",
    "Evidence_Path",
    "Notes",
}
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")


FORBIDDEN_ARTIFACT_SUFFIXES = {
    ".7z", ".a", ".bin", ".ckpt", ".dll", ".dylib", ".exe",
    ".gguf", ".gz", ".lib", ".onnx", ".pt", ".pth", ".pyd",
    ".safetensors", ".so", ".tar", ".tgz", ".whl", ".xml", ".zip",
}

ARTIFACT_CANONICAL_SOURCES: tuple[Path, ...] = (
    *(Path(workbook.template_path) for workbook in POST_C_WORKBOOKS),
    EXECUTION_INDEX_PATH,
    CONFIGURATION_MATRIX_PATH,
    EVIDENCE_REGISTER_PATH,
    REVISION_REGISTER_PATH,
    PACK_MANIFEST_PATH,
    CONTROLLED_MANIFEST_PATH,
)


def _issue(code: str, path: Path | str, message: str) -> PackIssue:
    """Create a stable issue while normalising local path separators."""

    return PackIssue(code, str(path).replace("\\", "/"), message)


def _sha256(path: Path) -> str:
    """Return the lowercase SHA-256 digest for one file."""

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """Read one UTF-8 CSV into an explicit header and row list."""

    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def _safe_relative_evidence_path(value: str) -> bool:
    """Accept empty evidence paths or normal repository-relative POSIX paths."""

    stripped = value.strip()
    if not stripped:
        return True

    # Reject Windows drive, UNC, device, absolute and parent-traversing forms.
    windows = PureWindowsPath(stripped)
    posix = PurePosixPath(stripped.replace("\\", "/"))
    if windows.drive or windows.root or posix.is_absolute():
        return False
    if any(part in {"", ".", ".."} for part in posix.parts):
        return False
    return not stripped.startswith(("~", "//", "\\\\"))


def _markdown_table_count(text: str) -> int:
    """Count canonical Markdown tables by their separator rows."""

    return sum(
        1
        for line in text.splitlines()
        if line.strip().startswith("|")
        and line.strip().endswith("|")
        and all(
            re.fullmatch(r":?-{3,}:?", cell.strip())
            for cell in line.strip()[1:-1].split("|")
        )
    )


def _validate_pack_manifest(root: Path, issues: list[PackIssue]) -> dict[str, object] | None:
    """Validate the controlling JSON manifest and return it when readable."""

    path = root / PACK_MANIFEST_PATH
    if not path.is_file():
        issues.append(_issue("PACK_MANIFEST_MISSING", PACK_MANIFEST_PATH, "Post-C pack manifest is missing."))
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        issues.append(_issue("PACK_MANIFEST_INVALID", PACK_MANIFEST_PATH, str(error)))
        return None

    if payload.get("pack_status") != "Initialised":
        issues.append(_issue("PACK_STATUS_INVALID", PACK_MANIFEST_PATH, "Initial pack status must be Initialised."))
    if payload.get("scientific_results_authorised") is not False:
        issues.append(_issue("SCIENTIFIC_AUTHORITY_INVALID", PACK_MANIFEST_PATH, "Initial pack must not authorise scientific results."))
    if payload.get("workbook_closure_authorised") is not False:
        issues.append(_issue("CLOSURE_AUTHORITY_INVALID", PACK_MANIFEST_PATH, "Initial pack must not authorise workbook closure."))
    if payload.get("stage_dependencies") != EXPECTED_STAGE_DEPENDENCIES:
        issues.append(_issue("STAGE_DEPENDENCIES_INVALID", PACK_MANIFEST_PATH, "Stage dependencies do not match the approved C-to-F order."))
    if set(payload.get("allowed_statuses", [])) != EXPECTED_STATUSES:
        issues.append(_issue("ALLOWED_STATUSES_INVALID", PACK_MANIFEST_PATH, "Allowed statuses do not match the controlled vocabulary."))

    entries = payload.get("workbooks")
    if not isinstance(entries, list) or len(entries) != len(POST_C_WORKBOOKS):
        issues.append(_issue("WORKBOOK_MANIFEST_COUNT_INVALID", PACK_MANIFEST_PATH, "Exactly seven post-C workbook entries are required."))
    else:
        expected_ids = {item.workbook_id for item in POST_C_WORKBOOKS}
        actual_ids = {str(entry.get("workbook_id")) for entry in entries if isinstance(entry, dict)}
        if actual_ids != expected_ids:
            issues.append(_issue("WORKBOOK_MANIFEST_IDS_INVALID", PACK_MANIFEST_PATH, "Workbook manifest IDs differ from WB-07 through WB-13."))
        for entry in entries:
            if isinstance(entry, dict) and entry.get("live_authority") is not False:
                issues.append(_issue("WORKBOOK_LIVE_AUTHORITY_INVALID", PACK_MANIFEST_PATH, f"{entry.get('workbook_id')} must begin with live_authority=false."))
    return payload


def _validate_workbooks(
    root: Path,
    generated_root: Path | None,
    manifest_payload: Mapping[str, object] | None,
    issues: list[PackIssue],
) -> None:
    """Validate every canonical Markdown workbook and generated DOCX."""

    manifest_entries: dict[str, Mapping[str, object]] = {}
    if manifest_payload and isinstance(manifest_payload.get("workbooks"), list):
        manifest_entries = {
            str(entry.get("workbook_id")): entry
            for entry in manifest_payload["workbooks"]
            if isinstance(entry, dict)
        }

    for workbook in POST_C_WORKBOOKS:
        template_relative = Path(workbook.template_path)
        generated_relative = Path(workbook.generated_path)
        template_path = root / template_relative
        generated_path = (
            generated_root / generated_relative.name
            if generated_root is not None
            else root / generated_relative
        )

        if not template_path.is_file():
            issues.append(_issue("TEMPLATE_MISSING", template_relative, f"{workbook.workbook_id} canonical Markdown is missing."))
            continue

        text = template_path.read_text(encoding="utf-8")
        if text.startswith("\ufeff"):
            issues.append(_issue("TEMPLATE_BOM", template_relative, "Canonical Markdown must be BOM-free UTF-8."))
        if "\r" in text:
            issues.append(_issue("TEMPLATE_LINE_ENDING", template_relative, "Canonical Markdown must use LF line endings."))
        for heading in workbook.required_headings:
            if not re.search(rf"^#+\s+{re.escape(heading)}\s*$", text, flags=re.MULTILINE | re.IGNORECASE):
                issues.append(_issue("REQUIRED_HEADING_MISSING", template_relative, f"Required heading is missing: {heading}"))
        if _markdown_table_count(text) < workbook.minimum_table_count:
            issues.append(_issue("TEMPLATE_TABLE_COUNT_LOW", template_relative, f"Expected at least {workbook.minimum_table_count} controlled tables."))
        if f"**Workbook ID:** `{workbook.workbook_id}`" not in text:
            issues.append(_issue("WORKBOOK_ID_MISSING", template_relative, f"Visible workbook identity {workbook.workbook_id} is missing."))
        if f"**Stage:** `{workbook.stage_id}`" not in text:
            issues.append(_issue("STAGE_ID_MISSING", template_relative, f"Visible stage identity {workbook.stage_id} is missing."))
        if "Status:** Initialised" not in text:
            issues.append(_issue("INITIAL_STATUS_MISSING", template_relative, "Workbook must visibly state Initialised status."))

        entry = manifest_entries.get(workbook.workbook_id)
        if entry:
            if entry.get("template_path") != workbook.template_path:
                issues.append(_issue("PACK_TEMPLATE_PATH_MISMATCH", PACK_MANIFEST_PATH, f"{workbook.workbook_id} template path drifted."))
            if entry.get("generated_path") != workbook.generated_path:
                issues.append(_issue("PACK_GENERATED_PATH_MISMATCH", PACK_MANIFEST_PATH, f"{workbook.workbook_id} generated path drifted."))

        if not generated_path.is_file():
            issues.append(_issue("DOCX_MISSING", generated_relative, f"{workbook.workbook_id} generated DOCX is missing."))
            continue
        if not zipfile.is_zipfile(generated_path):
            issues.append(_issue("DOCX_INVALID_PACKAGE", generated_relative, "Generated workbook is not a valid DOCX ZIP package."))
            continue
        try:
            document = Document(generated_path)
        except Exception as error:  # python-docx exposes multiple package errors.
            issues.append(_issue("DOCX_OPEN_FAILED", generated_relative, str(error)))
            continue

        visible_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        if workbook.workbook_id not in visible_text:
            issues.append(_issue("DOCX_WORKBOOK_ID_MISSING", generated_relative, f"Generated DOCX does not visibly contain {workbook.workbook_id}."))
        if len(document.tables) < workbook.minimum_table_count:
            issues.append(_issue("DOCX_TABLE_COUNT_LOW", generated_relative, f"Generated DOCX has {len(document.tables)} tables; expected at least {workbook.minimum_table_count}."))
        if not any(section.orientation is not None for section in document.sections):
            issues.append(_issue("DOCX_SECTION_MISSING", generated_relative, "Generated DOCX contains no readable section metadata."))


def _validate_controlled_manifest(
    root: Path,
    generated_root: Path | None,
    issues: list[PackIssue],
) -> None:
    """Verify the unified WB-01 through WB-13 manifest and new file hashes."""

    path = root / CONTROLLED_MANIFEST_PATH
    if not path.is_file():
        issues.append(_issue("CONTROLLED_MANIFEST_MISSING", CONTROLLED_MANIFEST_PATH, "Controlled workbook manifest is missing."))
        return
    try:
        headers, rows = _read_csv(path)
    except OSError as error:
        issues.append(_issue("CONTROLLED_MANIFEST_INVALID", CONTROLLED_MANIFEST_PATH, str(error)))
        return

    required = {
        "Workbook_ID",
        "Controlled_File",
        "Canonical_Text_Template",
        "Canonical_Template_SHA256",
        "Last_Validated_DOCX_SHA256",
        "Status",
        "Generation_Command",
        "Revision",
    }
    missing = required - set(headers)
    if missing:
        issues.append(_issue("CONTROLLED_MANIFEST_HEADERS_INVALID", CONTROLLED_MANIFEST_PATH, f"Missing headers: {sorted(missing)}"))
        return

    ids = [row["Workbook_ID"] for row in rows]
    if len(ids) != len(set(ids)):
        issues.append(_issue("DUPLICATE_WORKBOOK_ID", CONTROLLED_MANIFEST_PATH, "Workbook IDs must be unique."))
    if set(ids) != {f"WB-{number:02d}" for number in range(1, 14)}:
        issues.append(_issue("CONTROLLED_MANIFEST_IDS_INVALID", CONTROLLED_MANIFEST_PATH, "Controlled manifest must contain exactly WB-01 through WB-13."))

    by_id = {row["Workbook_ID"]: row for row in rows}
    for workbook in POST_C_WORKBOOKS:
        row = by_id.get(workbook.workbook_id)
        if row is None:
            continue
        template_path = root / workbook.template_path
        generated_path = (
            generated_root / Path(workbook.generated_path).name
            if generated_root is not None
            else root / workbook.generated_path
        )
        if row["Canonical_Text_Template"] != workbook.template_path:
            issues.append(_issue("CONTROLLED_TEMPLATE_PATH_MISMATCH", CONTROLLED_MANIFEST_PATH, f"{workbook.workbook_id} canonical path drifted."))
        if row["Controlled_File"] != Path(workbook.generated_path).name:
            issues.append(_issue("CONTROLLED_FILE_MISMATCH", CONTROLLED_MANIFEST_PATH, f"{workbook.workbook_id} controlled filename drifted."))
        if template_path.is_file() and row["Canonical_Template_SHA256"] != _sha256(template_path):
            issues.append(_issue("TEMPLATE_HASH_MISMATCH", CONTROLLED_MANIFEST_PATH, f"{workbook.workbook_id} template digest does not match the canonical file."))
        if generated_path.is_file() and row["Last_Validated_DOCX_SHA256"] != _sha256(generated_path):
            issues.append(_issue("DOCX_HASH_MISMATCH", CONTROLLED_MANIFEST_PATH, f"{workbook.workbook_id} DOCX digest does not match the generated file."))
        if not SHA256_RE.fullmatch(row["Canonical_Template_SHA256"]):
            issues.append(_issue("TEMPLATE_HASH_FORMAT_INVALID", CONTROLLED_MANIFEST_PATH, f"{workbook.workbook_id} template digest is not lowercase SHA-256."))
        if not SHA256_RE.fullmatch(row["Last_Validated_DOCX_SHA256"]):
            issues.append(_issue("DOCX_HASH_FORMAT_INVALID", CONTROLLED_MANIFEST_PATH, f"{workbook.workbook_id} DOCX digest is not lowercase SHA-256."))
        if row["Revision"] != "1.0":
            issues.append(_issue("POST_C_REVISION_INVALID", CONTROLLED_MANIFEST_PATH, f"{workbook.workbook_id} initial revision must be 1.0."))


def _validate_execution_index(root: Path, issues: list[PackIssue]) -> None:
    """Validate stage coverage, identities, statuses, paths and scientific boundaries."""

    path = root / EXECUTION_INDEX_PATH
    if not path.is_file():
        issues.append(_issue("EXECUTION_INDEX_MISSING", EXECUTION_INDEX_PATH, "Post-C execution index is missing."))
        return
    try:
        headers, rows = _read_csv(path)
    except OSError as error:
        issues.append(_issue("EXECUTION_INDEX_INVALID", EXECUTION_INDEX_PATH, str(error)))
        return

    missing = EXPECTED_INDEX_HEADERS - set(headers)
    if missing:
        issues.append(_issue("EXECUTION_INDEX_HEADERS_INVALID", EXECUTION_INDEX_PATH, f"Missing headers: {sorted(missing)}"))
        return
    if not rows:
        issues.append(_issue("EXECUTION_INDEX_EMPTY", EXECUTION_INDEX_PATH, "Execution index contains no records."))
        return

    record_ids = [row["Execution_Record_ID"] for row in rows]
    duplicate_ids = sorted(key for key, count in Counter(record_ids).items() if count > 1)
    for duplicate in duplicate_ids:
        issues.append(_issue("DUPLICATE_EXECUTION_RECORD", EXECUTION_INDEX_PATH, f"Execution record is duplicated: {duplicate}"))

    test_keys = [(row["Stage_ID"], row["Test_ID"], row["Repetition_Role"], row["Context_Tokens"]) for row in rows]
    duplicate_keys = sorted(key for key, count in Counter(test_keys).items() if count > 1)
    for duplicate in duplicate_keys:
        issues.append(_issue("DUPLICATE_EXECUTION_IDENTITY", EXECUTION_INDEX_PATH, f"Stage/test/repetition/context identity is duplicated: {duplicate}"))

    stages = {row["Stage_ID"] for row in rows}
    if stages != EXPECTED_STAGES:
        issues.append(_issue("STAGE_COVERAGE_INVALID", EXECUTION_INDEX_PATH, f"Observed stage set is {sorted(stages)}."))
    if sum(row["Stage_ID"] == "D2" for row in rows) != 48:
        issues.append(_issue("D2_ROW_COUNT_INVALID", EXECUTION_INDEX_PATH, "D2 must contain 12 Route A and 36 Route B ordered pairs."))
    if not any(row["Model_Scale"] == "3B" for row in rows):
        issues.append(_issue("GRANITE_3B_COVERAGE_MISSING", EXECUTION_INDEX_PATH, "No Granite 3B records exist."))
    if not any(row["Model_Scale"] == "8B" for row in rows):
        issues.append(_issue("GRANITE_8B_COVERAGE_MISSING", EXECUTION_INDEX_PATH, "No Granite 8B records exist."))
    if not any(row["Model_Scale"] == "30B" for row in rows):
        issues.append(_issue("GRANITE_30B_COVERAGE_MISSING", EXECUTION_INDEX_PATH, "No Granite 30B records exist."))

    for line_number, row in enumerate(rows, start=2):
        location = f"{EXECUTION_INDEX_PATH}:{line_number}"
        if row["Status"] not in EXPECTED_STATUSES:
            issues.append(_issue("STATUS_INVALID", location, f"Unsupported status: {row['Status']}"))
        if row["Frontier_Status"] not in EXPECTED_STATUSES:
            issues.append(_issue("FRONTIER_STATUS_INVALID", location, f"Unsupported frontier status: {row['Frontier_Status']}"))
        if row["Formal_Statistics_Allowed"] not in {"true", "false"}:
            issues.append(_issue("FORMAL_STATISTICS_FLAG_INVALID", location, "Formal_Statistics_Allowed must be true or false."))
        if row["Discovery_Only"] not in {"true", "false"}:
            issues.append(_issue("DISCOVERY_FLAG_INVALID", location, "Discovery_Only must be true or false."))
        if row["Stage_ID"] == "D2" and (row["Execution_Order"] or row["Memory_Rank"]):
            issues.append(_issue("PREMATURE_D2_ORDER", location, "D2 order and rank must remain blank until D1 verified storage is accepted."))
        if row["Stage_ID"] == "E3":
            if row["Formal_Statistics_Allowed"] != "false":
                issues.append(_issue("E3_FORMAL_AUTHORITY_INVALID", location, "30B feasibility cannot authorise formal statistics."))
            if row["Discovery_Only"] != "true":
                issues.append(_issue("E3_DISCOVERY_BOUNDARY_INVALID", location, "30B feasibility rows must remain discovery-only."))
        if row["Stage_ID"] == "F" and row["Status"] in {"Passed", "Not applicable"}:
            issues.append(_issue("PREMATURE_CLOSURE_STATUS", location, "F-stage rows cannot begin passed or not applicable."))
        if row["Run_ID"] or row["Artifact_ID"]:
            issues.append(_issue("PREMATURE_LIVE_IDENTITY", location, "Initial workbooks must not contain a live run or artifact identity."))
        if not _safe_relative_evidence_path(row["Evidence_Path"]):
            issues.append(_issue("UNSAFE_EVIDENCE_PATH", location, f"Unsafe evidence path: {row['Evidence_Path']}"))
        if row["Status"] == "Blocked pending predecessor" and not row["Blocker_Code"]:
            issues.append(_issue("BLOCKER_CODE_MISSING", location, "Predecessor-blocked rows require a stable blocker code."))


def _validate_shared_csvs(root: Path, issues: list[PackIssue]) -> None:
    """Validate configuration, evidence and revision registers."""

    # Configuration matrix controls route labels and keeps planning bytes distinct.
    config_path = root / CONFIGURATION_MATRIX_PATH
    if not config_path.is_file():
        issues.append(_issue("CONFIGURATION_MATRIX_MISSING", CONFIGURATION_MATRIX_PATH, "Configuration matrix is missing."))
    else:
        headers, rows = _read_csv(config_path)
        required = {
            "Configuration_ID",
            "Route_ID",
            "Role",
            "Model_Scales",
            "K_Algorithm",
            "K_Precision",
            "V_Algorithm",
            "V_Precision",
            "Admission_Status",
            "Evidence_Path",
        }
        if required - set(headers):
            issues.append(_issue("CONFIGURATION_HEADERS_INVALID", CONFIGURATION_MATRIX_PATH, f"Missing headers: {sorted(required - set(headers))}"))
        ids = [row["Configuration_ID"] for row in rows]
        if len(ids) != len(set(ids)):
            issues.append(_issue("DUPLICATE_CONFIGURATION_ID", CONFIGURATION_MATRIX_PATH, "Configuration IDs must be unique."))
        if not {"route-a-merged-openvino", "route-b-experimental-qjl-polar"}.issubset({row["Route_ID"] for row in rows}):
            issues.append(_issue("ROUTE_COVERAGE_INVALID", CONFIGURATION_MATRIX_PATH, "Route A and Route B controls are both required."))
        for line_number, row in enumerate(rows, start=2):
            if row["Evidence_Path"] and not _safe_relative_evidence_path(row["Evidence_Path"]):
                issues.append(_issue("UNSAFE_EVIDENCE_PATH", f"{CONFIGURATION_MATRIX_PATH}:{line_number}", f"Unsafe evidence path: {row['Evidence_Path']}"))
            if row["Route_ID"] == "route-a-merged-openvino" and any(token in row["Configuration_ID"] for token in ("QJL", "POLAR")):
                issues.append(_issue("ROUTE_LABEL_CONTRADICTION", f"{CONFIGURATION_MATRIX_PATH}:{line_number}", "QJL/Polar configuration cannot be labelled Route A merged OpenVINO."))

    # Evidence intake is append-only and must initially contain one empty row per stage.
    evidence_path = root / EVIDENCE_REGISTER_PATH
    if not evidence_path.is_file():
        issues.append(_issue("EVIDENCE_REGISTER_MISSING", EVIDENCE_REGISTER_PATH, "Evidence register is missing."))
    else:
        headers, rows = _read_csv(evidence_path)
        required = {
            "Evidence_Record_ID",
            "Stage_ID",
            "Artifact_SHA256",
            "Repository_Head_SHA",
            "Independent_Validation_Status",
            "Evidence_Path",
            "Ingestion_Status",
        }
        if required - set(headers):
            issues.append(_issue("EVIDENCE_HEADERS_INVALID", EVIDENCE_REGISTER_PATH, f"Missing headers: {sorted(required - set(headers))}"))
        if {row["Stage_ID"] for row in rows} != EXPECTED_STAGES:
            issues.append(_issue("EVIDENCE_STAGE_COVERAGE_INVALID", EVIDENCE_REGISTER_PATH, "Evidence register must initialise one row for D1 through F."))
        if any(row["Ingestion_Status"] != "Not started" for row in rows):
            issues.append(_issue("PREMATURE_EVIDENCE_INGESTION", EVIDENCE_REGISTER_PATH, "No evidence may be ingested in the initial structure."))
        for line_number, row in enumerate(rows, start=2):
            if row["Evidence_Path"] and not _safe_relative_evidence_path(row["Evidence_Path"]):
                issues.append(_issue("UNSAFE_EVIDENCE_PATH", f"{EVIDENCE_REGISTER_PATH}:{line_number}", f"Unsafe evidence path: {row['Evidence_Path']}"))

    # Revision register must preserve exactly one current initial row per workbook.
    revision_path = root / REVISION_REGISTER_PATH
    if not revision_path.is_file():
        issues.append(_issue("REVISION_REGISTER_MISSING", REVISION_REGISTER_PATH, "Post-C revision register is missing."))
    else:
        headers, rows = _read_csv(revision_path)
        required = {"Record_ID", "Workbook_ID", "Version", "Status", "Supersedes"}
        if required - set(headers):
            issues.append(_issue("REVISION_HEADERS_INVALID", REVISION_REGISTER_PATH, f"Missing headers: {sorted(required - set(headers))}"))
        expected_ids = {workbook.workbook_id for workbook in POST_C_WORKBOOKS}
        if {row["Workbook_ID"] for row in rows} != expected_ids:
            issues.append(_issue("REVISION_WORKBOOK_COVERAGE_INVALID", REVISION_REGISTER_PATH, "Revision register must cover WB-07 through WB-13."))
        for workbook_id in expected_ids:
            current = [row for row in rows if row["Workbook_ID"] == workbook_id and row["Status"].startswith("Current")]
            if len(current) != 1:
                issues.append(_issue("REVISION_CURRENT_COUNT_INVALID", REVISION_REGISTER_PATH, f"{workbook_id} must have exactly one current row."))
            elif current[0]["Version"] != "1.0":
                issues.append(_issue("REVISION_VERSION_INVALID", REVISION_REGISTER_PATH, f"{workbook_id} initial version must be 1.0."))


def _validate_artifact_bundle(
    repository_root: Path,
    artifact_root: Path,
    issues: list[PackIssue],
) -> None:
    """Validate the uploaded workbook artifact without trusting its catalogue."""

    if not artifact_root.is_dir():
        issues.append(_issue("ARTIFACT_ROOT_INVALID", artifact_root, "Artifact root is not a directory."))
        return
    manifest_path = artifact_root / "artifact-manifest.sha256"
    if not manifest_path.is_file():
        issues.append(_issue("ARTIFACT_MANIFEST_MISSING", manifest_path, "Artifact digest catalogue is missing."))
        return

    recorded: dict[str, str] = {}
    for line_number, line in enumerate(manifest_path.read_text(encoding="utf-8").splitlines(), start=1):
        match = re.fullmatch(r"([0-9a-f]{64})  (.+)", line)
        if match is None:
            issues.append(_issue("ARTIFACT_MANIFEST_LINE_INVALID", f"{manifest_path}:{line_number}", "Expected lowercase SHA-256, two spaces, and a relative path."))
            continue
        digest, relative = match.groups()
        if not _safe_relative_evidence_path(relative):
            issues.append(_issue("UNSAFE_ARTIFACT_PATH", f"{manifest_path}:{line_number}", f"Unsafe artifact path: {relative}"))
            continue
        if relative in recorded:
            issues.append(_issue("DUPLICATE_ARTIFACT_PATH", f"{manifest_path}:{line_number}", f"Duplicate artifact path: {relative}"))
            continue
        recorded[relative] = digest

    actual_files = {
        path.relative_to(artifact_root).as_posix(): path
        for path in artifact_root.rglob("*")
        if path.is_file() and path != manifest_path
    }
    if set(recorded) != set(actual_files):
        missing = sorted(set(recorded) - set(actual_files))
        extra = sorted(set(actual_files) - set(recorded))
        issues.append(_issue("ARTIFACT_FILE_SET_MISMATCH", artifact_root, f"Missing files: {missing}; unrecorded files: {extra}"))

    for relative, path in actual_files.items():
        if path.suffix.lower() in FORBIDDEN_ARTIFACT_SUFFIXES:
            issues.append(_issue("FORBIDDEN_ARTIFACT_PAYLOAD", relative, f"Forbidden artifact suffix: {path.suffix.lower()}"))
        expected = recorded.get(relative)
        if expected and _sha256(path) != expected:
            issues.append(_issue("ARTIFACT_HASH_MISMATCH", relative, "Artifact file digest differs from artifact-manifest.sha256."))

    # Canonical copies in the artifact must be byte-identical to the exact
    # checked-out validation sources. Basenames are unique within this pack.
    canonical_root = artifact_root / "canonical"
    for source_relative in ARTIFACT_CANONICAL_SOURCES:
        source = repository_root / source_relative
        captured = canonical_root / source_relative.name
        if not captured.is_file():
            issues.append(_issue("ARTIFACT_CANONICAL_FILE_MISSING", captured, f"Missing captured canonical file for {source_relative}."))
        elif source.is_file() and _sha256(captured) != _sha256(source):
            issues.append(_issue("ARTIFACT_CANONICAL_HASH_MISMATCH", captured, f"Captured canonical file differs from {source_relative}."))


def validate_post_c_workbook_pack(
    repository_root: Path,
    generated_directory: Path | None = None,
    artifact_root: Path | None = None,
) -> list[PackIssue]:
    """Return all stable validation issues for one post-C workbook pack."""

    root = repository_root.resolve()
    issues: list[PackIssue] = []
    if not root.is_dir():
        return [_issue("REPOSITORY_ROOT_INVALID", root, "Repository root is not a directory.")]

    manifest_payload = _validate_pack_manifest(root, issues)
    _validate_workbooks(root, generated_directory.resolve() if generated_directory else None, manifest_payload, issues)
    _validate_controlled_manifest(
        root,
        generated_directory.resolve() if generated_directory else None,
        issues,
    )
    _validate_execution_index(root, issues)
    _validate_shared_csvs(root, issues)
    if artifact_root is not None:
        _validate_artifact_bundle(root, artifact_root.resolve(), issues)
    return sorted(issues, key=lambda item: (item.code, item.path, item.message))


def _write_report(path: Path, issues: Sequence[PackIssue]) -> None:
    """Write one deterministic JSON validation report."""

    payload = {
        "schema_version": "1.0",
        "status": "Passed" if not issues else "Failed",
        "issue_count": len(issues),
        "issues": [
            {"code": issue.code, "path": issue.path, "message": issue.message}
            for issue in issues
        ],
        "scientific_results_authorised": False,
        "workbook_closure_authorised": False,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8", newline="\n")


def main(argv: Iterable[str] | None = None) -> int:
    """Validate a pack from the command line and return a process exit code."""

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repository-root", type=Path, required=True)
    parser.add_argument("--generated-directory", type=Path)
    parser.add_argument("--artifact-root", type=Path)
    parser.add_argument("--report", type=Path)
    arguments = parser.parse_args(list(argv) if argv is not None else None)

    generated_directory = arguments.generated_directory
    if arguments.artifact_root is not None and generated_directory is None:
        generated_directory = arguments.artifact_root / "generated"
    issues = validate_post_c_workbook_pack(
        arguments.repository_root,
        generated_directory,
        arguments.artifact_root,
    )
    if arguments.report:
        _write_report(arguments.report, issues)
    if issues:
        for issue in issues:
            print(f"{issue.code}: {issue.path}: {issue.message}", file=sys.stderr)
        return 1
    print("WORKBOOK05_POST_C_WORKBOOK_PACK_VALID")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
