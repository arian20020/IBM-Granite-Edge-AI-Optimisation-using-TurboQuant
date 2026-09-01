"""Structural fallback QA when a compatible DOCX renderer is unavailable."""

from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

from docx import Document


REQUIRED_TEXT = ("AH-B01", "AH-B08", "AH-01", "AH-10", "6.23", "3.66",
                 "6.39", "6.73", "3.28", "6.27", "zero unresolved failures",
                 "Completed fresh CPU, SYCL and supplementary Vulkan builds",
                 "Current - pending merge", "1.3")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    with zipfile.ZipFile(args.docx) as package:
        bad_member = package.testzip()
        member_count = len(package.namelist())
    if bad_member:
        raise ValueError(f"corrupt DOCX member: {bad_member}")
    document = Document(args.docx)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    text = "\n".join(paragraphs + cells)
    missing = [value for value in REQUIRED_TEXT if value not in text]
    if missing:
        raise ValueError(f"DOCX required content missing: {missing}")
    blank_cells = sum(1 for value in cells if not value.strip())
    if blank_cells:
        raise ValueError(f"DOCX contains {blank_cells} blank table cells")
    result = {"docx": str(args.docx), "zip_members": member_count,
              "paragraphs": len(paragraphs), "tables": len(document.tables),
              "table_cells": len(cells), "blank_table_cells": 0,
              "required_text_checks": len(REQUIRED_TEXT),
              "visual_render": "not completed: LibreOffice absent and hidden Word PDF export exceeded 120 seconds"}
    args.output.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
