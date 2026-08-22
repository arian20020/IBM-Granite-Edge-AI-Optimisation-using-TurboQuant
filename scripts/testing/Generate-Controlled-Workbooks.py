"""Generate thirteen controlled DOCX testing workbooks from canonical Markdown.

The Markdown files are the canonical, reviewable workbook templates in Git. The
DOCX files are deterministic working/reporting artefacts that can be opened and
completed in Microsoft Word. ``--post-c-only`` limits generation to WB-07 through
WB-13 so the repository-safe post-C workflow does not rewrite legacy workbooks.
"""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LEGACY_WORKBOOK_TEMPLATES: Sequence[str] = (
    "01_Upstream_llama.cpp_Controlled_Retest_Workbook_v1.md",
    "02_AtomicBot_TurboQuant_Controlled_Retest_Workbook_v1.md",
    "03_animehacker_TQ3_0_Controlled_Retest_Workbook_v1.md",
    "04_Official_OpenVINO_Controlled_Retest_Workbook_v1.md",
    "05_Custom_OpenVINO_TurboQuant_Controlled_Retest_Workbook_v1.md",
    "06_Cross_Route_Controlled_Comparison_Workbook_v1.md",
)

POST_C_WORKBOOK_TEMPLATES: Sequence[str] = (
    "07_Workbook_05_D1_Codec_Conformance_Controlled_Workbook_v1.md",
    "08_Workbook_05_D2_Diagnostic_KV_Sweep_Controlled_Workbook_v1.md",
    "09_Workbook_05_E1_Granite_3B_Frontier_Formal_Controlled_Workbook_v1.md",
    "10_Workbook_05_E2_Granite_8B_Feasibility_Controlled_Workbook_v1.md",
    "11_Workbook_05_E3_Granite_30B_Bounded_Feasibility_Controlled_Workbook_v1.md",
    "12_Workbook_05_E4_Cross_Family_Repeatability_Controlled_Workbook_v1.md",
    "13_Workbook_05_F_Evidence_Closure_Controlled_Workbook_v1.md",
)

WORKBOOK_TEMPLATES: Sequence[str] = (
    *LEGACY_WORKBOOK_TEMPLATES,
    *POST_C_WORKBOOK_TEMPLATES,
)


def clean_cell(value: str) -> str:
    """Remove Markdown escaping used inside table cells."""

    return value.strip().replace("\\|", "|").replace("<br>", "\n")


def is_separator_row(cells: Sequence[str]) -> bool:
    """Return True when every table cell contains Markdown separator syntax."""

    return bool(cells) and all(
        re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells
    )


def parse_table(
    lines: Sequence[str],
    start_index: int,
) -> tuple[List[List[str]], int]:
    """Read consecutive Markdown table rows and return rows plus next index."""

    raw_rows: List[List[str]] = []
    index = start_index
    while index < len(lines):
        line = lines[index].strip()
        if not (line.startswith("|") and line.endswith("|")):
            break
        raw_rows.append(
            [clean_cell(cell) for cell in line[1:-1].split("|")]
        )
        index += 1
    if len(raw_rows) >= 2 and is_separator_row(raw_rows[1]):
        raw_rows.pop(1)
    width = max((len(row) for row in raw_rows), default=0)
    return [row + [""] * (width - len(row)) for row in raw_rows], index


def set_cell_fill(cell, fill_hex: str) -> None:
    """Apply a solid hexadecimal fill colour to a Word table cell."""

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill_hex)


def mark_repeat_header(row) -> None:
    """Repeat one table header row at the top of later pages."""

    row_properties = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    row_properties.append(table_header)


def prevent_row_split(row) -> None:
    """Keep a logical table row together on a page whenever Word can."""

    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))


def style_table(table) -> None:
    """Format a generated table for clear printing and Word editing."""

    table.style = "Table Grid"
    table.autofit = True
    for row in table.rows:
        prevent_row_split(row)
    if table.rows:
        mark_repeat_header(table.rows[0])
        for cell in table.rows[0].cells:
            set_cell_fill(cell, "1F4E78")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(7)
    for row_index, row in enumerate(table.rows[1:], start=1):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                set_cell_fill(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(7)


def configure_document(document: Document, footer_text: str) -> None:
    """Apply the controlled workbook's page and typography settings."""

    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = (
        section.page_height,
        section.page_width,
    )
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    for style_name, size in (
        ("Title", 19),
        ("Heading 1", 15),
        ("Heading 2", 12),
        ("Heading 3", 10),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 78, 120)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(footer_text)
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)


def add_inline_markdown(paragraph, text: str) -> None:
    """Convert paired double-asterisk spans into bold Word runs."""

    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_text_paragraph(document: Document, text: str) -> None:
    """Add one non-table Markdown line to the Word document."""

    if text.startswith("- "):
        paragraph = document.add_paragraph(style="List Bullet")
        content = text[2:]
    elif re.match(r"^\d+\.\s+", text):
        paragraph = document.add_paragraph(style="List Number")
        content = re.sub(r"^\d+\.\s+", "", text)
    else:
        paragraph = document.add_paragraph()
        content = text
    add_inline_markdown(paragraph, content)
    paragraph.paragraph_format.space_after = Pt(3)


def normalize_docx_package(output_path: Path) -> None:
    """Rewrite the DOCX ZIP with stable member ordering and timestamps."""

    with zipfile.ZipFile(output_path, "r") as source:
        members = [
            (info, source.read(info.filename))
            for info in source.infolist()
        ]
    temporary_path = output_path.with_suffix(output_path.suffix + ".tmp")
    with zipfile.ZipFile(
        temporary_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as target:
        for original, data in sorted(
            members,
            key=lambda item: item[0].filename,
        ):
            info = zipfile.ZipInfo(
                original.filename,
                date_time=(1980, 1, 1, 0, 0, 0),
            )
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.internal_attr = original.internal_attr
            info.create_system = original.create_system
            info.flag_bits = original.flag_bits
            target.writestr(info, data)
    temporary_path.replace(output_path)


def convert_markdown_to_docx(
    markdown_path: Path,
    output_path: Path,
) -> None:
    """Create one deterministic Word workbook from canonical Markdown."""

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    document = Document()

    # Preserve the historic generator metadata so WB-01 through WB-06 remain
    # byte-compatible when regenerated with the extended template list.
    fixed_timestamp = datetime(2026, 7, 13, 0, 0, 0)
    document.core_properties.author = (
        "Granite-TurboQuant Controlled Retest Campaign"
    )
    document.core_properties.last_modified_by = (
        "Granite-TurboQuant Controlled Retest Campaign"
    )
    document.core_properties.created = fixed_timestamp
    document.core_properties.modified = fixed_timestamp
    document.core_properties.revision = 1
    configure_document(
        document,
        "Granite-TurboQuant Controlled Retest Campaign v1 - generated working copy",
    )

    index = 0
    first_heading_used = False
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped == "[[PAGEBREAK]]":
            document.add_page_break()
            index += 1
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            heading_text = heading_match.group(2).strip()
            if not first_heading_used:
                paragraph = document.add_paragraph(
                    heading_text,
                    style="Title",
                )
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_heading_used = True
            else:
                document.add_heading(heading_text, level=min(level, 3))
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            rows, next_index = parse_table(lines, index)
            if rows and rows[0]:
                table = document.add_table(
                    rows=len(rows),
                    cols=len(rows[0]),
                )
                for row_index, row_values in enumerate(rows):
                    for column_index, value in enumerate(row_values):
                        table.cell(row_index, column_index).text = value
                style_table(table)
                document.add_paragraph().paragraph_format.space_after = Pt(2)
            index = next_index
            continue

        add_text_paragraph(document, stripped)
        index += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    normalize_docx_package(output_path)


def default_repository_root() -> Path:
    """Return the repository root while this script remains in scripts/testing."""

    return Path(__file__).resolve().parents[2]


def main(argv: Iterable[str] | None = None) -> int:
    """Generate the selected workbook package and return an exit code."""

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--repository-root",
        type=Path,
        default=default_repository_root(),
    )
    parser.add_argument("--output-directory", type=Path, default=None)
    parser.add_argument(
        "--post-c-only",
        action="store_true",
        help="Generate only WB-07 through WB-13.",
    )
    arguments = parser.parse_args(list(argv) if argv is not None else None)

    repository_root = arguments.repository_root.resolve()
    template_directory = (
        repository_root
        / "docs"
        / "testing"
        / "workbooks"
        / "text-templates"
    )
    output_directory = (
        arguments.output_directory.resolve()
        if arguments.output_directory is not None
        else repository_root
        / "docs"
        / "testing"
        / "workbooks"
        / "generated"
    )
    selected_templates = (
        POST_C_WORKBOOK_TEMPLATES
        if arguments.post_c_only
        else WORKBOOK_TEMPLATES
    )

    if not template_directory.is_dir():
        print(
            f"ERROR: Template directory not found: {template_directory}",
            file=sys.stderr,
        )
        return 2
    for template_name in selected_templates:
        template_path = template_directory / template_name
        if not template_path.is_file():
            print(
                f"ERROR: Required template not found: {template_path}",
                file=sys.stderr,
            )
            return 3
        output_path = output_directory / template_path.with_suffix(".docx").name
        convert_markdown_to_docx(template_path, output_path)
        print(f"Generated: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
