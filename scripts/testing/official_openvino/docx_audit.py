"""Structural and control audit for the generated WB-04 DOCX."""

from __future__ import annotations

import csv
import hashlib
import math
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document

from scripts.testing.finalize_official_openvino_comparison_workbook import (
    COMPARISON_SECTION_TITLES,
)
from scripts.testing.official_openvino.matrix import COMPARISON_IDS as MATRIX_COMPARISON_IDS


PRESENTATION_MEASURED_PAIRS = (
    ("OV-TQ-13", 512),
    ("OV-TQ-14", 512),
    ("OV-TQ-14", 2048),
)
PRESENTATION_LIMITATION_BULLETS = (
    "Diagnostic only; no formal benchmark",
    "Missing validated FP16 artifact",
    "RAM safety floor reached",
    "Larger host required",
    "Strict activation proof incomplete",
    "Governed quality campaign stopped at the RAM floor",
)
NO_SCORE_NO_WINNER_DISCLOSURE = (
    "No numeric quality score exists and there is no winner."
)
PROHIBITED_PRESENTATION_CLAIMS = (
    "all tests passed",
    "quality-qualified pass",
    "quality winner",
)

REVISION_HISTORY_HEADERS = (
    "Version", "Date", "Changed by", "Change", "Affected test IDs",
    "Reference", "Status",
)
CURRENT_REVISION_ID = "WR-036"
CURRENT_REVISION_VERSION = "1.8"
CURRENT_REVISION_DATE = "2026-07-31"
CURRENT_REVISION_STATUS = "Current - pending merge"


@dataclass(frozen=True)
class DocxAuditProfile:
    workbook_version: str
    revision_id: str
    revision_date: str
    current_status: str
    expected_ids: frozenset[str]
    cache_ids: frozenset[str]
    comparison_headings: tuple[str, ...]
    expected_matrix_sha256: str


HISTORICAL_AUDIT_PROFILE = DocxAuditProfile(
    workbook_version=CURRENT_REVISION_VERSION,
    revision_id=CURRENT_REVISION_ID,
    revision_date=CURRENT_REVISION_DATE,
    current_status=CURRENT_REVISION_STATUS,
    expected_ids=frozenset(),
    cache_ids=frozenset(),
    comparison_headings=(),
    expected_matrix_sha256=(
        "7db2636b403d285aa3886c9f23560a9e48bd43645e9aca35e0d1fc4f16eaea42"
    ),
)


def comparison_audit_profile(matrix_sha256: str) -> DocxAuditProfile:
    if (
        not isinstance(matrix_sha256, str)
        or not re.fullmatch(r"[0-9a-f]{64}", matrix_sha256)
    ):
        raise ValueError("comparison matrix must be a lowercase SHA-256")
    return DocxAuditProfile(
        workbook_version="1.9",
        revision_id="WR-037",
        revision_date="2026-08-01",
        current_status="Current - pending PR",
        expected_ids=frozenset(MATRIX_COMPARISON_IDS),
        cache_ids=frozenset({"OV-12", "OV-TQ-21", "OV-TQ-22"}),
        comparison_headings=COMPARISON_SECTION_TITLES,
        expected_matrix_sha256=matrix_sha256,
    )


def _presentation_pair_is_visible(
    visible_text: str, test_id: str, context_tokens: int
) -> bool:
    return re.search(
        rf"(?<![A-Z0-9-]){re.escape(test_id)}(?![A-Z0-9-])"
        rf"\s*(?:\|\s*)?{context_tokens}\b",
        visible_text,
    ) is not None


def audit_presentation(visible_text: str) -> dict[str, int]:
    """Validate the compact v1.8 presentation without relaxing control checks."""
    missing_pairs = [
        f"{test_id}/{context_tokens}"
        for test_id, context_tokens in PRESENTATION_MEASURED_PAIRS
        if not _presentation_pair_is_visible(visible_text, test_id, context_tokens)
    ]
    if missing_pairs:
        raise ValueError(
            "missing measured presentation pairs: " + ", ".join(missing_pairs)
        )
    missing_limitations = [
        bullet
        for bullet in PRESENTATION_LIMITATION_BULLETS
        if bullet not in visible_text
    ]
    if missing_limitations:
        raise ValueError(
            "missing compact limitation bullets: " + ", ".join(missing_limitations)
        )
    if NO_SCORE_NO_WINNER_DISCLOSURE not in visible_text:
        raise ValueError("exact no-score/no-winner disclosure is missing")

    lowered = visible_text.casefold()
    for claim in PROHIBITED_PRESENTATION_CLAIMS:
        if claim in lowered:
            raise ValueError(f"prohibited presentation claim: {claim}")
    quality_score_search_text = re.sub(r"\s+", " ", visible_text)
    if re.search(
        r"\bquality[ -]score\b(?:\s*[:=|]\s*|\s+)"
        r"(?<![a-z0-9])[0-9]+(?:\.[0-9]+)?(?:\s*/\s*10)?\b",
        quality_score_search_text,
        re.IGNORECASE,
    ):
        raise ValueError("prohibited numeric quality score")
    without_no_winner = re.sub(
        r"\b(?:there\s+is\s+)?no(?:\s+[a-z-]+){0,3}\s+winner\b",
        "",
        visible_text,
        flags=re.IGNORECASE,
    )
    if re.search(r"\bwinner\b", without_no_winner, re.IGNORECASE):
        raise ValueError("prohibited winner claim")
    return {
        "presentation_measured_row_count": len(PRESENTATION_MEASURED_PAIRS),
        "presentation_limitation_bullet_count": len(PRESENTATION_LIMITATION_BULLETS),
    }


def audit_revision_history(
    document: Document,
    visible_text: str,
    profile: DocxAuditProfile = HISTORICAL_AUDIT_PROFILE,
) -> dict[str, int]:
    """Validate the unique revision-history table and exact current release row."""

    if profile.revision_id not in visible_text:
        raise ValueError(f"visible {profile.revision_id} revision is missing")
    history_tables = [
        table
        for table in document.tables
        if tuple(cell.text.strip() for cell in table.rows[0].cells)
        == REVISION_HISTORY_HEADERS
    ]
    if len(history_tables) != 1:
        raise ValueError(
            "DOCX must contain exactly one revision-history table with exact headers"
        )
    table = history_tables[0]
    rows = [
        {
            header: cell.text.strip()
            for header, cell in zip(REVISION_HISTORY_HEADERS, row.cells)
        }
        for row in table.rows[1:]
    ]
    current = [row for row in rows if row["Status"].startswith("Current")]
    if len(current) != 1:
        raise ValueError("revision history must contain exactly one current row")
    row = current[0]
    for field, expected in (
        ("Version", profile.workbook_version),
        ("Date", profile.revision_date),
        ("Status", profile.current_status),
    ):
        if row[field] != expected:
            raise ValueError(
                f"revision-history current {field} must be {expected}; got {row[field]}"
            )
    return {"revision_history_current_rows": 1}


def _zip_integrity(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            bad_zip_member = archive.testzip()
    except zipfile.BadZipFile as error:
        raise ValueError("DOCX audit found an invalid ZIP package") from error
    if bad_zip_member is not None:
        raise ValueError(f"corrupt DOCX member: {bad_zip_member}")


def _visible_document(document: Document) -> tuple[str, list[list[list[str]]], list[list[int]]]:
    text_parts = [paragraph.text for paragraph in document.paragraphs]
    tables: list[list[list[str]]] = []
    blank_cells: list[list[int]] = []
    for table_index, table in enumerate(document.tables):
        rows: list[list[str]] = []
        for row_index, row in enumerate(table.rows):
            values: list[str] = []
            for cell_index, cell in enumerate(row.cells):
                value = cell.text.strip()
                values.append(value)
                text_parts.append(cell.text)
                if not value:
                    blank_cells.append([table_index, row_index, cell_index])
            rows.append(values)
        tables.append(rows)
    return "\n".join(text_parts), tables, blank_cells


def _split_markdown_row(line: str) -> list[str]:
    stripped = line.strip()
    if not (stripped.startswith("|") and stripped.endswith("|")):
        raise ValueError("comparison Markdown table row is malformed")
    cells: list[str] = []
    current: list[str] = []
    escaped = False
    for character in stripped[1:-1]:
        if escaped:
            current.append(character)
            escaped = False
        elif character == "\\":
            escaped = True
        elif character == "|":
            cells.append("".join(current).strip().replace("<br>", "\n"))
            current = []
        else:
            current.append(character)
    if escaped:
        current.append("\\")
    cells.append("".join(current).strip().replace("<br>", "\n"))
    return cells


def _markdown_tables(text: str) -> list[list[list[str]]]:
    lines = text.splitlines()
    result: list[list[list[str]]] = []
    index = 0
    while index < len(lines):
        if not lines[index].strip().startswith("|"):
            index += 1
            continue
        raw: list[list[str]] = []
        while index < len(lines) and lines[index].strip().startswith("|"):
            raw.append(_split_markdown_row(lines[index]))
            index += 1
        if len(raw) < 2 or not all(
            re.fullmatch(r":?-{3,}:?", cell) for cell in raw[1]
        ):
            raise ValueError("comparison Markdown table separator is invalid")
        width = len(raw[0])
        if any(len(row) != width for row in raw):
            raise ValueError("comparison Markdown table row width differs")
        result.append([raw[0], *raw[2:]])
    return result


def _manifest_row(path: Path) -> dict[str, str]:
    with Path(path).open(newline="", encoding="utf-8-sig") as handle:
        rows = [row for row in csv.DictReader(handle) if row.get("Workbook_ID") == "WB-04"]
    if len(rows) != 1:
        raise ValueError("controlled manifest must contain exactly one WB-04 row")
    return rows[0]


_CACHE_TABLE_HEADER = (
    "Test ID",
    "Context",
    "Cache route",
    "Activation",
    "U8 artifact SHA-256",
    "Runtime evidence",
)
_QUALITY_TABLE_HEADER = (
    "Test ID",
    "Context",
    "P1",
    "P2",
    "P3",
    "P4",
    "P5",
    "P6",
    "Mean",
    "Median",
    "Minimum",
    "Maximum",
    "Quality evidence",
)


def _winner_from_staged_tables(
    tables: list[list[list[str]]],
    winner_lines: tuple[str, ...],
    profile: DocxAuditProfile,
) -> str | None:
    if len(winner_lines) > 1:
        raise ValueError("comparison winner claim is duplicated")
    cache_matches = [
        table for table in tables if table and tuple(table[0]) == _CACHE_TABLE_HEADER
    ]
    quality_matches = [
        table for table in tables if table and tuple(table[0]) == _QUALITY_TABLE_HEADER
    ]
    if not cache_matches and not quality_matches and not winner_lines:
        return None
    if len(cache_matches) != 1 or len(quality_matches) != 1:
        if winner_lines:
            raise ValueError("comparison winner lacks exact eligibility tables")
        return None
    cache = cache_matches[0]
    quality = quality_matches[0]
    cache_ids = set(profile.cache_ids)
    cache_pairs: set[tuple[str, str]] = set()
    contexts: set[str] = set()
    for row in cache[1:]:
        test_id, context = row[0], row[1]
        if test_id not in cache_ids or not context.isdigit():
            raise ValueError("comparison winner cache rows are not exact")
        pair = (test_id, context)
        if pair in cache_pairs:
            raise ValueError("comparison winner cache row is duplicated")
        cache_pairs.add(pair)
        contexts.add(context)
    if not contexts or cache_pairs != {
        (test_id, context) for context in contexts for test_id in cache_ids
    }:
        if winner_lines:
            raise ValueError("comparison winner lacks full shared cache context coverage")
        return None

    quality_by_pair: dict[tuple[str, str], list[float]] = {}
    for row in quality[1:]:
        pair = (row[0], row[1])
        if pair not in cache_pairs:
            continue
        if pair in quality_by_pair:
            raise ValueError("comparison winner quality row is duplicated")
        values: list[float] = []
        for raw in row[2:12]:
            try:
                number = float(raw)
            except ValueError as error:
                raise ValueError("comparison winner quality row is incomplete") from error
            if not math.isfinite(number):
                raise ValueError("comparison winner quality row is non-finite")
            values.append(number)
        quality_by_pair[pair] = values[:6]
    if set(quality_by_pair) != cache_pairs:
        if winner_lines:
            raise ValueError("comparison winner lacks complete numeric quality rows")
        return None
    scores = {
        test_id: [
            score
            for context in contexts
            for score in quality_by_pair[(test_id, context)]
        ]
        for test_id in cache_ids
    }
    means = {
        test_id: sum(values) / len(values) for test_id, values in scores.items()
    }
    best = max(means.values())
    eligible = [test_id for test_id, value in means.items() if value == best]
    if len(eligible) != 1:
        if winner_lines:
            raise ValueError("comparison winner is tied or unsupported")
        return None
    winner = eligible[0]
    if not winner_lines:
        raise ValueError("comparison winner is missing despite unique eligibility")
    expected = (
        f"Overall winner: {winner} (shared complete numeric quality)."
    )
    if winner_lines != (expected,):
        raise ValueError("comparison winner claim differs from recomputed eligibility")
    return winner


def audit_comparison_docx(
    docx_path: Path,
    manifest_path: Path,
    markdown_path: Path,
    profile: DocxAuditProfile,
) -> dict[str, object]:
    """Audit the staged v1.9 DOCX against its finalized Markdown source."""

    if profile.workbook_version != "1.9" or profile.revision_id != "WR-037":
        raise ValueError("comparison DOCX audit requires the v1.9 profile")
    _zip_integrity(Path(docx_path))
    markdown = Path(markdown_path).read_text(encoding="utf-8-sig")
    document = Document(docx_path)
    visible_text, actual_tables, blank_cells = _visible_document(document)
    if blank_cells:
        raise ValueError(f"blank DOCX table cells: {blank_cells[:5]}")
    missing_ids = sorted(
        test_id for test_id in profile.expected_ids if test_id not in visible_text
    )
    if missing_ids:
        raise ValueError(f"missing controlled IDs in DOCX: {missing_ids}")
    paragraph_text = [paragraph.text.strip() for paragraph in document.paragraphs]
    for number, heading in enumerate(profile.comparison_headings, start=5):
        visible_heading = f"{number}. {heading}"
        if paragraph_text.count(visible_heading) != 1:
            raise ValueError(
                f"comparison heading is missing or duplicated: {visible_heading}"
            )
        if markdown.count(f"# {visible_heading}") != 1:
            raise ValueError(f"staged Markdown comparison heading is invalid: {heading}")
    matrix_sha = profile.expected_matrix_sha256
    if markdown.count(matrix_sha) != 1 or visible_text.count(matrix_sha) != 1:
        raise ValueError("comparison matrix SHA-256 is missing or duplicated")
    if HISTORICAL_AUDIT_PROFILE.expected_matrix_sha256 in markdown or (
        HISTORICAL_AUDIT_PROFILE.expected_matrix_sha256 in visible_text
    ):
        raise ValueError("historical matrix SHA-256 remains in the comparison release")
    revision = audit_revision_history(document, visible_text, profile)
    revision_tables = [
        index
        for index, table in enumerate(actual_tables)
        if table and tuple(table[0]) == REVISION_HISTORY_HEADERS
    ]
    if len(revision_tables) != 1:
        raise ValueError("DOCX must contain exactly one revision-history table")
    content_tables = [
        table for index, table in enumerate(actual_tables) if index != revision_tables[0]
    ]
    expected_tables = _markdown_tables(markdown)
    if content_tables != expected_tables:
        raise ValueError("DOCX comparison tables differ from finalized Markdown")
    markdown_winners = tuple(
        re.sub(r"\*\*", "", line.strip())
        for line in markdown.splitlines()
        if line.strip().startswith("Overall winner:")
    )
    visible_winners = tuple(
        line.strip()
        for line in paragraph_text
        if line.strip().startswith("Overall winner:")
    )
    if visible_winners != markdown_winners:
        raise ValueError("unsupported winner claim in DOCX")
    comparison_winner = _winner_from_staged_tables(
        expected_tables, markdown_winners, profile
    )
    manifest = _manifest_row(manifest_path)
    if manifest.get("Revision") != profile.workbook_version:
        raise ValueError("WB-04 manifest revision differs from comparison profile")
    actual_hash = hashlib.sha256(Path(docx_path).read_bytes()).hexdigest()
    if manifest.get("Last_Validated_DOCX_SHA256", "").lower() != actual_hash:
        raise ValueError("WB-04 generated hash differs from controlled manifest")
    return {
        "accepted": True,
        "zip_integrity": "passed",
        "controlled_id_count": len(profile.expected_ids),
        "blank_table_cells": 0,
        "visible_revision": profile.workbook_version,
        "revision_history": "validated",
        "generated_sha256": actual_hash,
        "table_count": len(actual_tables),
        "comparison_heading_count": len(profile.comparison_headings),
        "comparison_table_count": len(content_tables),
        "comparison_winner": comparison_winner,
        **revision,
    }


def audit_docx(docx_path: Path, manifest_path: Path, required_ids: set[str]) -> dict:
    _zip_integrity(Path(docx_path))

    document = Document(docx_path)
    visible_text, _tables, blank_cells = _visible_document(document)
    missing_ids = sorted(test_id for test_id in required_ids if test_id not in visible_text)
    if blank_cells:
        raise ValueError(f"blank DOCX table cells: {blank_cells[:5]}")
    if missing_ids:
        raise ValueError(f"missing controlled IDs in DOCX: {missing_ids}")
    presentation = audit_presentation(visible_text)
    revision_history = audit_revision_history(
        document, visible_text, HISTORICAL_AUDIT_PROFILE
    )
    with manifest_path.open(newline="", encoding="utf-8-sig") as handle:
        manifest_row = next(row for row in csv.DictReader(handle)
                            if row["Workbook_ID"] == "WB-04")
    manifest_revision = manifest_row["Revision"].strip()
    if not manifest_revision:
        raise ValueError("WB-04 manifest revision is missing")
    visible_revision = f"v{manifest_revision}"
    if visible_revision not in visible_text:
        raise ValueError(
            f"visible {visible_revision} title is missing"
        )

    actual_hash = hashlib.sha256(docx_path.read_bytes()).hexdigest()
    if actual_hash.lower() != manifest_row["Last_Validated_DOCX_SHA256"].lower():
        raise ValueError("WB-04 generated hash differs from controlled manifest")
    return {
        "accepted": True,
        "zip_integrity": "passed",
        "controlled_id_count": len(required_ids),
        "blank_table_cells": 0,
        "visible_revision": manifest_revision,
        "revision_history": "validated",
        "generated_sha256": actual_hash,
        "table_count": len(document.tables),
        **presentation,
        **revision_history,
    }
