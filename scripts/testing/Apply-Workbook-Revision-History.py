"""Apply the Git-controlled revision history to the six generated testing workbooks."""
from __future__ import annotations

import argparse
import csv
import re
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

WORKBOOKS = {
    "01_Upstream_llama.cpp_Controlled_Retest_Workbook_v1.docx": "WB-01",
    "02_AtomicBot_TurboQuant_Controlled_Retest_Workbook_v1.docx": "WB-02",
    "03_animehacker_TQ3_0_Controlled_Retest_Workbook_v1.docx": "WB-03",
    "04_Official_OpenVINO_Controlled_Retest_Workbook_v1.docx": "WB-04",
    "05_Custom_OpenVINO_TurboQuant_Controlled_Retest_Workbook_v1.docx": "WB-05",
    "06_Cross_Route_Controlled_Comparison_Workbook_v1.docx": "WB-06",
}


def select_workbooks(workbook_id: str | None) -> dict[str, str]:
    """Return all workbooks, or the one controlled workbook explicitly selected."""
    if workbook_id is None:
        return WORKBOOKS
    selected = {
        filename: candidate_id
        for filename, candidate_id in WORKBOOKS.items()
        if candidate_id == workbook_id
    }
    if not selected:
        raise ValueError(f"Unknown workbook ID: {workbook_id}")
    return selected


def repository_root() -> Path:
    return Path(__file__).resolve().parents[2]


def load_history(path: Path) -> dict[str, list[dict[str, str]]]:
    grouped = {workbook_id: [] for workbook_id in WORKBOOKS.values()}
    with path.open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            grouped[row["Workbook_ID"]].append(row)
    for workbook_id, rows in grouped.items():
        if not rows:
            raise ValueError(f"No revision history found for {workbook_id}")
        rows.sort(key=lambda row: tuple(int(part) for part in row["Version"].split(".")))
        if sum(row["Status"].startswith("Current") for row in rows) != 1:
            raise ValueError(f"{workbook_id} must have exactly one current revision")
    return grouped


def set_fill(cell, value: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), value)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    properties.append(element)


def style_table(table) -> None:
    table.style = "Table Grid"
    table.autofit = True
    repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        prevent_row_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_fill(cell, "1F4E78")
            elif row_index % 2 == 0:
                set_fill(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(7)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def normalize_package(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]
    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for original, data in sorted(members, key=lambda item: item[0].filename):
            info = zipfile.ZipInfo(original.filename, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.internal_attr = original.internal_attr
            info.create_system = original.create_system
            info.flag_bits = original.flag_bits
            target.writestr(info, data)
    temporary.replace(path)


def update_visible_version(document: Document, version: str) -> None:
    title = document.paragraphs[0]
    title.text = re.sub(r"\s+v\d+(?:\.\d+)?$", f" v{version}", title.text)
    for paragraph in document.paragraphs[1:]:
        if paragraph.text.startswith("Controlled retest template v"):
            paragraph.text = re.sub(r"v\d+(?:\.\d+)?", f"v{version}", paragraph.text, count=1)
            break
        if paragraph.text.startswith("Controlled retest revision "):
            paragraph.text = re.sub(r"(Controlled retest revision )\d+(?:\.\d+)?", rf"\g<1>{version}", paragraph.text, count=1)
            break


def insert_history(document: Document, rows: list[dict[str, str]]) -> None:
    title = document.paragraphs[0]
    heading = document.add_heading("Document revision history", level=1)
    table = document.add_table(rows=1, cols=7)
    headers = ["Version", "Date", "Changed by", "Change", "Affected test IDs", "Reference", "Status"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        cells = table.add_row().cells
        reference = row["Change_Reference"]
        values = [row["Version"], row["Date"], row["Changed_By"], row["Change_Summary"], row["Affected_Test_IDs"], reference, row["Status"]]
        for cell, value in zip(cells, values):
            cell.text = value
    style_table(table)
    note = document.add_paragraph("Canonical source: docs/testing/Workbook-Revision-Register.csv. Older rows are retained and never overwritten.")
    note.runs[0].bold = True
    note.paragraph_format.space_after = Pt(4)
    # addnext inserts immediately after title, so move elements in reverse display order.
    title._p.addnext(note._p)
    title._p.addnext(table._tbl)
    title._p.addnext(heading._p)


def process(source: Path, destination: Path, rows: list[dict[str, str]]) -> None:
    document = Document(source)
    current = next(row for row in rows if row["Status"].startswith("Current"))
    update_visible_version(document, current["Version"])
    insert_history(document, rows)
    fixed = datetime(2026, 7, 14, 0, 0, 0)
    document.core_properties.modified = fixed
    document.core_properties.revision = 1
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    normalize_package(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repository-root", type=Path, default=repository_root())
    parser.add_argument("--input-directory", type=Path)
    parser.add_argument("--output-directory", type=Path)
    parser.add_argument(
        "--workbook-id",
        choices=tuple(WORKBOOKS.values()),
        help="Apply history only to one controlled workbook (for example WB-04).",
    )
    args = parser.parse_args()
    root = args.repository_root.resolve()
    input_dir = (args.input_directory or root / "docs/testing/workbooks/generated").resolve()
    output_dir = (args.output_directory or input_dir).resolve()
    histories = load_history(root / "docs/testing/Workbook-Revision-Register.csv")
    for filename, workbook_id in select_workbooks(args.workbook_id).items():
        source = input_dir / filename
        if not source.is_file():
            raise FileNotFoundError(source)
        destination = output_dir / filename
        process(source, destination, histories[workbook_id])
        print(f"Revision history applied: {destination}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
