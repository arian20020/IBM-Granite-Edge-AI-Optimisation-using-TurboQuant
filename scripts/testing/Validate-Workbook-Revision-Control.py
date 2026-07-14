"""Validate workbook revision, source, template and generated-DOCX controls."""

from __future__ import annotations

import csv
import hashlib
import re
import zipfile
from collections import defaultdict
from pathlib import Path

from docx import Document


# These are the only workbook identifiers controlled by this campaign.
WORKBOOK_IDS = {
    "WB-01",
    "WB-02",
    "WB-03",
    "WB-04",
    "WB-05",
    "WB-06",
}

# Versions must contain only dotted numeric components.
VERSION_RE = re.compile(r"^\d+(?:\.\d+)+$")

# Final references require a pull-request number and a full 40-character commit.
FINAL_REFERENCE_RE = re.compile(r"#\d+.*\b[0-9a-fA-F]{40}\b")

# Pre-merge references require Pending merge and either Pending PR or a PR number.
PENDING_REFERENCE_RE = re.compile(
    r"(?:Pending PR|#\d+).*Pending merge",
    re.IGNORECASE,
)

# All SHA-256 values must contain exactly 64 lowercase hexadecimal characters.
SHA256_RE = re.compile(r"^[0-9a-f]{64}$")


def sha256(path: Path) -> str:
    """Return the lowercase SHA-256 value of one file."""
    digest = hashlib.sha256()

    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)

    return digest.hexdigest()


def version_tuple(value: str) -> tuple[int, ...]:
    """Convert a dotted numeric version into a comparable integer tuple."""
    if not VERSION_RE.fullmatch(value):
        raise ValueError(f"Invalid version: {value}")

    return tuple(int(part) for part in value.split("."))


def valid_current_reference(reference: str) -> bool:
    """Accept either a valid pre-merge or final merged lifecycle reference."""
    if "Pending merge" in reference:
        return bool(PENDING_REFERENCE_RE.search(reference))

    return bool(FINAL_REFERENCE_RE.search(reference))


def main() -> int:
    """Run every workbook-control check and return a process status."""
    root = Path(__file__).resolve().parents[2]
    errors: list[str] = []

    register_path = root / "docs/testing/Workbook-Revision-Register.csv"
    manifest_path = root / "docs/testing/workbooks/Controlled-Workbook-Manifest.csv"
    source_manifest_path = root / "docs/testing/source-material/Source-Document-Manifest.csv"
    generated_directory = root / "docs/testing/workbooks/generated"

    # Read the complete append-only revision history.
    with register_path.open(newline="", encoding="utf-8-sig") as handle:
        revision_rows = list(csv.DictReader(handle))

    # Read the six controlled workbook records.
    with manifest_path.open(newline="", encoding="utf-8-sig") as handle:
        manifest_rows = list(csv.DictReader(handle))

    # Read the original source-document register.
    with source_manifest_path.open(newline="", encoding="utf-8-sig") as handle:
        source_rows = list(csv.DictReader(handle))

    # Record identifiers must be unique and sequential.
    record_ids = [row["Record_ID"] for row in revision_rows]

    if len(record_ids) != len(set(record_ids)):
        errors.append("Revision register contains duplicate Record_ID values")

    numeric_record_ids: list[int] = []

    for record_id in record_ids:
        match = re.fullmatch(r"WR-(\d{3})", record_id)

        if not match:
            errors.append(f"Invalid revision Record_ID: {record_id}")
            continue

        numeric_record_ids.append(int(match.group(1)))

    if numeric_record_ids and sorted(numeric_record_ids) != list(
        range(1, len(numeric_record_ids) + 1)
    ):
        errors.append("Revision Record_ID sequence is not continuous")

    # Group all revision records by workbook.
    histories: defaultdict[str, list[dict[str, str]]] = defaultdict(list)

    for row in revision_rows:
        histories[row["Workbook_ID"]].append(row)

    if set(histories) != WORKBOOK_IDS:
        errors.append("Workbook IDs do not match WB-01 to WB-06")

    # Index the controlled workbook manifest by ID.
    manifest_by_id = {
        row["Workbook_ID"]: row
        for row in manifest_rows
    }

    if len(manifest_rows) != 6 or set(manifest_by_id) != WORKBOOK_IDS:
        errors.append("Workbook manifest must contain one row for WB-01 to WB-06")

    # Index source-document rows by source filename.
    source_rows_by_name: defaultdict[str, list[dict[str, str]]] = defaultdict(list)

    for source_row in source_rows:
        source_name = Path(source_row["Source_Path"]).name
        source_rows_by_name[source_name].append(source_row)

    for workbook_id in sorted(WORKBOOK_IDS):
        history = histories.get(workbook_id, [])

        try:
            history = sorted(
                history,
                key=lambda row: version_tuple(row["Version"]),
            )
        except ValueError as exc:
            errors.append(f"{workbook_id}: {exc}")
            continue

        current_rows = [
            row
            for row in history
            if row["Status"].startswith("Current")
        ]

        if len(current_rows) != 1:
            errors.append(
                f"{workbook_id}: expected exactly one current revision row"
            )
            continue

        current_row = current_rows[0]

        # Every older record must remain present and explicitly superseded.
        for row in history:
            if row is current_row:
                continue

            if row["Status"] != "Superseded":
                errors.append(
                    f"{workbook_id}: older revision {row['Version']} "
                    f"is not Superseded"
                )

        if history[-1] is not current_row:
            errors.append(
                f"{workbook_id}: current revision is not the highest version"
            )

        if not valid_current_reference(current_row["Change_Reference"]):
            errors.append(
                f"{workbook_id}: current PR/merge lifecycle reference is invalid"
            )

        manifest_row = manifest_by_id.get(workbook_id)

        if manifest_row is None:
            errors.append(f"{workbook_id}: manifest row missing")
            continue

        if manifest_row["Revision"] != current_row["Version"]:
            errors.append(
                f"{workbook_id}: manifest version does not match current revision"
            )

        # Validate all three controlled SHA-256 fields.
        for field in (
            "Source_SHA256",
            "Canonical_Template_SHA256",
            "Last_Validated_DOCX_SHA256",
        ):
            if not SHA256_RE.fullmatch(manifest_row.get(field, "")):
                errors.append(
                    f"{workbook_id}: invalid manifest {field}"
                )

        # The controlled source hash must agree with the source-document manifest.
        source_name = manifest_row["Source_File"]
        matching_source_rows = source_rows_by_name.get(source_name, [])

        if len(matching_source_rows) != 1:
            errors.append(
                f"{workbook_id}: expected one source-manifest row for {source_name}"
            )
        else:
            source_row = matching_source_rows[0]
            source_hash = source_row["Original_SHA256"]

            if not SHA256_RE.fullmatch(source_hash):
                errors.append(
                    f"{workbook_id}: source-document manifest hash is invalid"
                )
            elif manifest_row["Source_SHA256"] != source_hash:
                errors.append(
                    f"{workbook_id}: controlled Source_SHA256 does not match "
                    "Source-Document-Manifest.csv"
                )

        template_path = root / manifest_row["Canonical_Text_Template"]
        docx_path = generated_directory / manifest_row["Controlled_File"]

        # Validate the canonical template and its exact content controls.
        if not template_path.is_file():
            errors.append(f"{workbook_id}: canonical template missing")
        else:
            if sha256(template_path) != manifest_row["Canonical_Template_SHA256"]:
                errors.append(f"{workbook_id}: template hash mismatch")

            template_text = template_path.read_text(encoding="utf-8")

            required_metadata_lines = (
                f"**Controlled filename:** `{manifest_row['Controlled_File']}`",
                "**Generated DOCX hash:** recorded in "
                "`Controlled-Workbook-Manifest.csv`",
                f"**Original source:** `{manifest_row['Source_File']}`",
                f"**Original source SHA-256:** "
                f"`{manifest_row['Source_SHA256']}`",
            )

            for required_line in required_metadata_lines:
                if required_line not in template_text:
                    errors.append(
                        f"{workbook_id}: canonical template metadata missing: "
                        f"{required_line}"
                    )

            # Legacy output hashes are forbidden because they become stale and
            # can be mistaken for the current generated workbook's own hash.
            if "**Source file:**" in template_text:
                errors.append(
                    f"{workbook_id}: legacy Source file metadata remains"
                )

            if re.search(r"(?m)^\*\*SHA-256:\*\*", template_text):
                errors.append(
                    f"{workbook_id}: legacy embedded generated hash remains"
                )

        # Validate the generated Word package and its recorded hash.
        if not docx_path.is_file():
            errors.append(f"{workbook_id}: generated DOCX missing")
            continue

        if sha256(docx_path) != manifest_row["Last_Validated_DOCX_SHA256"]:
            errors.append(f"{workbook_id}: generated DOCX hash mismatch")

        try:
            with zipfile.ZipFile(docx_path) as archive:
                bad_member = archive.testzip()

            if bad_member:
                errors.append(
                    f"{workbook_id}: corrupt DOCX member {bad_member}"
                )

            document = Document(docx_path)

            if not any(
                paragraph.text == "Document revision history"
                for paragraph in document.paragraphs
            ):
                errors.append(
                    f"{workbook_id}: revision heading missing"
                )

            history_tables = [
                table
                for table in document.tables
                if table.rows
                and table.cell(0, 0).text == "Version"
            ]

            if not history_tables:
                errors.append(
                    f"{workbook_id}: revision table missing"
                )
            elif len(history_tables[0].rows) - 1 != len(history):
                errors.append(
                    f"{workbook_id}: embedded history row count mismatch"
                )

            if not document.paragraphs:
                errors.append(
                    f"{workbook_id}: document contains no visible title"
                )
            elif not document.paragraphs[0].text.endswith(
                "v" + current_row["Version"]
            ):
                errors.append(
                    f"{workbook_id}: visible title version mismatch"
                )

        except Exception as exc:
            errors.append(
                f"{workbook_id}: DOCX validation failed: {exc}"
            )

    if errors:
        print("WORKBOOK REVISION CONTROL: FAIL")

        for error in errors:
            print(" -", error)

        return 1

    print("WORKBOOK REVISION CONTROL: PASS")
    print(
        f"{len(revision_rows)} append-only revision records, six current "
        "revisions, source-document hashes, canonical metadata and synchronized "
        "embedded histories validated."
    )
    print(
        "Boundary: this validates document change control, "
        "not hardware or model test success."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
