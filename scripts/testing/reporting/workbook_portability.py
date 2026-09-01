"""Create deterministic portable XLSX derivatives without changing source evidence."""

from __future__ import annotations

import hashlib
import os
import re
import tempfile
import zipfile
from pathlib import Path, PurePosixPath, PureWindowsPath

import openpyxl
from openpyxl import load_workbook

from .csvio import write_json


_PROJECT_DIRECTORY = "IBM-Granite-Edge-AI-Optimisation-using-TurboQuant"
_REPOSITORY_TOP_LEVELS = {
    ".github",
    ".superpowers",
    "docs",
    "experiments",
    "external",
    "IBM Granite with TurboQuant (Intel)",
    "models",
    "outputs",
    "release-evidence",
    "report",
    "research",
    "runtime",
    "scripts",
    "tests",
    "third-party",
}
_WINDOWS_ABSOLUTE = re.compile(r"^[A-Za-z]:[\\/]")
_FILE_URI = re.compile(r"^file:/+", re.IGNORECASE)
_PORTABLE_WORKBOOKS = (
    {
        "route": "04-openvino-experimental-fork",
        "activity_id": "#experimental-openvino-portable-workbook-generation",
        "source": (
            "results/source/"
            "Granite_OpenVINO_Final_Healthcare_Education_Results_2026-08-30.xlsx"
        ),
        "output": (
            "workbook/generated/"
            "openvino-experimental-fork-portable-results.xlsx"
        ),
        "receipt": (
            "workbook/generated/"
            "portable-workbook-provenance.json"
        ),
    },
    {
        "route": "05-openvino-official-upstream",
        "activity_id": "#official-openvino-portable-workbook-generation",
        "source": (
            "results/source/"
            "Granite_Official_OpenVINO_TurboQuant_Results_2026-08-30_v2_Missing_Attempts.xlsx"
        ),
        "output": (
            "workbook/generated/"
            "openvino-official-upstream-portable-results.xlsx"
        ),
        "receipt": (
            "workbook/generated/"
            "portable-workbook-provenance.json"
        ),
    },
)


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _is_machine_absolute(value: object) -> bool:
    if not isinstance(value, str):
        return False
    return bool(
        _WINDOWS_ABSOLUTE.match(value)
        or _FILE_URI.match(value)
        or value.startswith("\\\\")
    )


def _portable_reference(value: str) -> str:
    candidate = _FILE_URI.sub("", value).replace("/", "\\")
    parts = list(PureWindowsPath(candidate).parts)
    folded = [part.casefold() for part in parts]

    relative_parts: list[str] | None = None
    if ".worktrees" in folded:
        marker = folded.index(".worktrees")
        if len(parts) > marker + 2:
            relative_parts = parts[marker + 2 :]
    elif _PROJECT_DIRECTORY.casefold() in folded:
        marker = folded.index(_PROJECT_DIRECTORY.casefold())
        if len(parts) > marker + 1:
            relative_parts = parts[marker + 1 :]

    if not relative_parts:
        raise ValueError(
            "absolute workbook reference is not attributable to this repository"
        )
    if relative_parts[0] not in _REPOSITORY_TOP_LEVELS:
        raise ValueError(
            "absolute workbook reference does not resolve to a known repository root"
        )
    portable = PurePosixPath(*relative_parts)
    if portable.is_absolute() or ".." in portable.parts:
        raise ValueError("portable workbook reference escapes the repository")
    return portable.as_posix()


def _normalize_package(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=".p.", suffix=".tmp", dir=path.parent
    )
    try:
        os.close(descriptor)
        with zipfile.ZipFile(
            temporary_name, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
        ) as target:
            for original, data in sorted(members, key=lambda member: member[0].filename):
                info = zipfile.ZipInfo(original.filename, (1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = original.external_attr
                info.internal_attr = original.internal_attr
                info.create_system = original.create_system
                info.comment = original.comment
                target.writestr(info, data)
        Path(temporary_name).replace(path)
    finally:
        Path(temporary_name).unlink(missing_ok=True)


def _structural_signature(workbook) -> dict[str, object]:
    return {
        "sheetnames": tuple(workbook.sheetnames),
        "sheets": {
            sheet.title: {
                "dimensions": sheet.calculate_dimension(),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged": tuple(sorted(str(item) for item in sheet.merged_cells.ranges)),
                "freeze_panes": str(sheet.freeze_panes or ""),
                "auto_filter": str(sheet.auto_filter.ref or ""),
                "tables": tuple(
                    sorted((table.name, table.ref) for table in sheet.tables.values())
                ),
                "state": sheet.sheet_state,
            }
            for sheet in workbook.worksheets
        },
    }


def _validate_semantic_preservation(source: Path, portable: Path) -> None:
    original = load_workbook(source, data_only=False, keep_links=True)
    derivative = load_workbook(portable, data_only=False, keep_links=True)
    try:
        if _structural_signature(original) != _structural_signature(derivative):
            raise ValueError("portable workbook changed sheet structure")
        for source_sheet, portable_sheet in zip(
            original.worksheets, derivative.worksheets, strict=True
        ):
            for row in source_sheet.iter_rows():
                for source_cell in row:
                    portable_cell = portable_sheet[source_cell.coordinate]
                    expected_value = (
                        _portable_reference(source_cell.value)
                        if _is_machine_absolute(source_cell.value)
                        else source_cell.value
                    )
                    if portable_cell.value != expected_value:
                        raise ValueError(
                            "portable workbook changed a scientific value or formula: "
                            f"{source_sheet.title}!{source_cell.coordinate}"
                        )
                    source_target = (
                        source_cell.hyperlink.target
                        if source_cell.hyperlink is not None
                        else None
                    )
                    expected_target = (
                        _portable_reference(source_target)
                        if _is_machine_absolute(source_target)
                        else source_target
                    )
                    portable_target = (
                        portable_cell.hyperlink.target
                        if portable_cell.hyperlink is not None
                        else None
                    )
                    if portable_target != expected_target:
                        raise ValueError(
                            "portable workbook changed a hyperlink unexpectedly: "
                            f"{source_sheet.title}!{source_cell.coordinate}"
                        )
                    if _is_machine_absolute(portable_cell.value) or _is_machine_absolute(
                        portable_target
                    ):
                        raise ValueError(
                            "portable workbook retains an absolute machine path: "
                            f"{source_sheet.title}!{source_cell.coordinate}"
                        )
    finally:
        original.close()
        derivative.close()


def _formula_coordinates(path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=True)
    try:
        return [
            (sheet.title, cell.coordinate)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
    finally:
        workbook.close()


def _cached_formula_count(
    path: Path, coordinates: list[tuple[str, str]]
) -> int:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=True)
    try:
        return sum(
            workbook[sheet][coordinate].value is not None
            for sheet, coordinate in coordinates
        )
    finally:
        workbook.close()


def sanitize_workbook(source: Path, destination: Path) -> dict[str, object]:
    """Write a deterministic portable derivative and return its audit facts."""
    source_path = Path(source).resolve(strict=True)
    destination_path = Path(destination).resolve(strict=False)
    if source_path == destination_path:
        raise ValueError("portable workbook destination must differ from source evidence")

    source_sha256 = _hash_file(source_path)
    source_size = source_path.stat().st_size
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(source_path, data_only=False, keep_links=True)
    replacements: list[dict[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if _is_machine_absolute(cell.value):
                        cell.value = _portable_reference(cell.value)
                        replacements.append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "kind": "value",
                                "portable_value": cell.value,
                            }
                        )
                    if cell.hyperlink and _is_machine_absolute(cell.hyperlink.target):
                        cell.hyperlink.target = _portable_reference(cell.hyperlink.target)
                        replacements.append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "kind": "hyperlink",
                                "portable_value": cell.hyperlink.target,
                            }
                        )

        descriptor, temporary_name = tempfile.mkstemp(
            prefix=".p.",
            suffix=".xlsx",
            dir=destination_path.parent,
        )
        os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            workbook.save(temporary)
            _normalize_package(temporary)
            _validate_semantic_preservation(source_path, temporary)
            temporary.replace(destination_path)
        finally:
            temporary.unlink(missing_ok=True)
    finally:
        workbook.close()

    formulas = _formula_coordinates(destination_path)
    return {
        "schema_version": "1.0.0",
        "source_sha256": source_sha256,
        "source_size_bytes": source_size,
        "output_sha256": _hash_file(destination_path),
        "output_size_bytes": destination_path.stat().st_size,
        "replacement_count": len(replacements),
        "replacements": replacements,
        "formula_count": len(formulas),
        "cached_formula_value_count": _cached_formula_count(
            destination_path, formulas
        ),
        "openpyxl_version": openpyxl.__version__,
    }


def write_portable_openvino_workbooks(
    repository_root: Path,
) -> tuple[dict[str, object], ...]:
    """Publish both OpenVINO portable workbooks and provenance receipts."""
    root = Path(repository_root).resolve(strict=True)
    collection = root / "docs/testing/final-results"
    return tuple(
        write_portable_openvino_route_workbook(
            root, collection / str(configuration["route"])
        )
        for configuration in _PORTABLE_WORKBOOKS
    )


def write_portable_openvino_route_workbook(
    repository_root: Path, route: Path
) -> dict[str, object]:
    """Publish the portable XLSX and receipt for one OpenVINO route."""
    root = Path(repository_root).resolve(strict=True)
    route_path = Path(route).resolve(strict=True)
    route_path.relative_to(root)
    matches = [
        item for item in _PORTABLE_WORKBOOKS if item["route"] == route_path.name
    ]
    if len(matches) != 1:
        raise ValueError(f"unsupported OpenVINO route for workbook portability: {route_path.name}")
    configuration = matches[0]
    source = route_path / str(configuration["source"])
    output = route_path / str(configuration["output"])
    receipt_path = route_path / str(configuration["receipt"])
    facts = sanitize_workbook(source, output)
    receipt = {
        **facts,
        "activity_id": configuration["activity_id"],
        "source_path": source.relative_to(root).as_posix(),
        "source_role": "immutable evidence-only nonportable workbook",
        "output_path": output.relative_to(root).as_posix(),
        "output_role": "primary portable workbook derivative",
        "receipt_path": receipt_path.relative_to(root).as_posix(),
        "sanitizer": "scripts/testing/reporting/workbook_portability.py",
        "scientific_values_preserved": True,
        "formulas_preserved": True,
        "sheet_structure_preserved": True,
        "machine_absolute_path_count": 0,
    }
    write_json(receipt_path, receipt)
    return receipt


__all__ = [
    "sanitize_workbook",
    "write_portable_openvino_route_workbook",
    "write_portable_openvino_workbooks",
]
