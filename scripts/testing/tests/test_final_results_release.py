from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path, PurePosixPath
from urllib.parse import urlsplit

import fitz
import openpyxl


sys.path.insert(0, str(Path(__file__).resolve().parents[3]))


from scripts.testing.final_results.validate import (
    validate_collection,
    validate_release_metadata,
)


REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
RELEASE_ROOT = REPOSITORY_ROOT / "docs/testing/final-results"
ROUTES = (
    "01-upstream-llama-cpp",
    "02-atomicbot-turboquant",
    "03-animehacker-tq3-0",
    "04-openvino-experimental-fork",
    "05-openvino-official-upstream",
    "06-cross-route-comparison",
)
REPORT_STEMS = {
    "01-upstream-llama-cpp": "upstream-llama-cpp-final-report",
    "02-atomicbot-turboquant": "atomicbot-turboquant-final-report",
    "03-animehacker-tq3-0": "animehacker-tq3-0-final-report",
    "04-openvino-experimental-fork": "openvino-experimental-fork-final-report",
    "05-openvino-official-upstream": "openvino-official-upstream-final-report",
    "06-cross-route-comparison": "cross-route-comparison-final-report",
}
OPENVINO_WORKBOOKS = {
    "04-openvino-experimental-fork/results/source/Granite_OpenVINO_Final_Healthcare_Education_Results_2026-08-30.xlsx": (
        "Dashboard",
        "Format Comparison",
        "Sector Summary",
        "Detailed Results",
        "Quality Details",
        "Availability Matrix",
        "Methodology",
        "Source Data",
    ),
    "05-openvino-official-upstream/results/source/Granite_Official_OpenVINO_TurboQuant_Results_2026-08-30_v2_Missing_Attempts.xlsx": (
        "Executive Summary",
        "Format Comparison",
        "Sector Summary",
        "Detailed Results",
        "Quality Details",
        "Availability Matrix",
        "Missing Attempts",
        "Methodology",
        "Source Data",
    ),
}
CATALOGS = (
    "catalog/route-register.csv",
    "catalog/campaign-summary.csv",
    "catalog/performance-summary.csv",
    "catalog/quality-summary.csv",
    "catalog/failure-summary.csv",
    "catalog/evidence-manifest.csv",
    "catalog/claim-evidence-map.csv",
    "catalog/comparability-matrix.csv",
)
FORMULA_ERROR_VALUES = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
}


def _markdown_targets(path: Path) -> set[str]:
    text = path.read_text(encoding="utf-8")
    return {
        target.split("#", 1)[0]
        for target in re.findall(r"\[[^\]]+\]\(([^)]+)\)", text)
        if target and not urlsplit(target).scheme
    }


def _type_names(entity: dict[str, object]) -> set[str]:
    value = entity.get("@type", ())
    if isinstance(value, str):
        return {value}
    return {str(item) for item in value}


def _graph() -> tuple[dict[str, object], dict[str, dict[str, object]]]:
    crate = json.loads(
        (RELEASE_ROOT / "ro-crate-metadata.json").read_text(encoding="utf-8")
    )
    entities = {str(entity["@id"]): entity for entity in crate["@graph"]}
    return crate, entities


def test_release_portal_links_every_route_report_workbook_and_control_surface():
    targets = _markdown_targets(RELEASE_ROOT / "README.md")
    expected = {
        "CHANGELOG.md",
        "REPRODUCING.md",
        "LICENSES.md",
        "ro-crate-metadata.json",
        "manifest-sha256.txt",
        "validation/release-readiness.json",
        *CATALOGS,
    }
    for route in ROUTES:
        stem = REPORT_STEMS[route]
        expected.update(
            {
                f"{route}/route-manifest.json",
                f"{route}/workbook/source/{stem}.md",
                f"{route}/workbook/generated/{stem}.docx",
                f"{route}/workbook/generated/{stem}.pdf",
            }
        )
    expected.update(OPENVINO_WORKBOOKS)

    assert expected <= targets
    assert all((RELEASE_ROOT / target).is_file() for target in expected)


def test_ro_crate_is_version_13_and_all_packaged_data_entity_ids_are_relative():
    crate, entities = _graph()

    assert crate["@context"][0] == "https://w3id.org/ro/crate/1.3/context"
    descriptor = entities["ro-crate-metadata.json"]
    assert descriptor["conformsTo"] == {"@id": "https://w3id.org/ro/crate/1.3"}
    assert descriptor["about"] == {"@id": "./"}
    assert "Dataset" in _type_names(entities["./"])

    for entity_id, entity in entities.items():
        if not (_type_names(entity) & {"File", "Dataset"}):
            continue
        parsed = urlsplit(entity_id)
        assert not parsed.scheme, entity_id
        assert not parsed.netloc, entity_id
        assert not entity_id.startswith(("/", "\\")), entity_id
        assert not re.match(r"^[A-Za-z]:", entity_id), entity_id
        assert ".." not in PurePosixPath(entity_id).parts, entity_id


def test_ro_crate_covers_reports_tables_systems_and_generation_provenance():
    _, entities = _graph()
    expected_files = set(CATALOGS)
    for route in ROUTES:
        stem = REPORT_STEMS[route]
        expected_files.update(
            {
                f"{route}/workbook/source/{stem}.md",
                f"{route}/workbook/generated/{stem}.docx",
                f"{route}/workbook/generated/{stem}.pdf",
            }
        )
    for route in ROUTES[:5]:
        expected_files.update(
            {
                f"{route}/results/attempts.csv",
                f"{route}/results/measurements.csv",
                f"{route}/results/summary-results.csv",
            }
        )
    expected_files.update(OPENVINO_WORKBOOKS)
    assert expected_files <= entities.keys()

    for route in ROUTES[:5]:
        assert f"#{route}-repository" in entities
        assert f"#{route}-hardware" in entities
        assert f"#{route}-software" in entities

    generated = {
        entity_id: entity
        for entity_id, entity in entities.items()
        if "#generated-artifact" in {
            str(value.get("@id"))
            for value in entity.get("additionalType", ())
            if isinstance(value, dict)
        }
    }
    assert {
        f"{route}/workbook/generated/{REPORT_STEMS[route]}.{suffix}"
        for route in ROUTES
        for suffix in ("docx", "pdf")
    } <= generated.keys()
    assert "manifest-sha256.txt" in generated

    for entity_id, entity in generated.items():
        activity_id = entity["wasGeneratedBy"]["@id"]
        activity = entities[activity_id]
        assert "CreateAction" in _type_names(activity)
        assert {item["@id"] for item in activity["result"]} >= {entity_id}
        assert activity.get("object"), entity_id


def test_release_manifest_is_complete_sorted_self_excluding_and_hash_valid():
    manifest = RELEASE_ROOT / "manifest-sha256.txt"
    lines = manifest.read_text(encoding="utf-8").splitlines()
    entries: dict[str, str] = {}
    for line in lines:
        match = re.fullmatch(r"([0-9a-f]{64})  ([^\\]+)", line)
        assert match, line
        digest, relative = match.groups()
        assert relative not in entries
        entries[relative] = digest

    expected = {
        path.relative_to(RELEASE_ROOT).as_posix()
        for path in RELEASE_ROOT.rglob("*")
        if path.is_file() and path != manifest
    }
    assert list(entries) == sorted(entries)
    assert "manifest-sha256.txt" not in entries
    assert set(entries) == expected
    for relative, expected_digest in entries.items():
        actual = hashlib.sha256((RELEASE_ROOT / relative).read_bytes()).hexdigest()
        assert actual == expected_digest, relative


def test_licensing_gaps_are_explicit_and_citation_is_not_invented():
    text = (RELEASE_ROOT / "LICENSES.md").read_text(encoding="utf-8").casefold()

    assert "no repository-level license file" in text
    assert "no license grant" in text
    assert "not verified" in text
    assert all(route in text for route in ROUTES[:5])
    assert not (RELEASE_ROOT / "CITATION.cff").exists()
    assert "citation.cff" in text
    assert "verified author metadata" in text


def test_reproduction_commands_are_exact_and_preserve_the_no_rerun_boundary():
    text = (RELEASE_ROOT / "REPRODUCING.md").read_text(encoding="utf-8")

    assert "--validate-only" in text
    assert "test_final_results_*.py" in text
    assert "does not rerun benchmarks" in text.casefold()
    assert "does not modify raw evidence" in text.casefold()
    for route in ROUTES:
        assert f"{route}/reproduction/" in text


def test_openvino_workbooks_pass_read_only_sheet_and_formula_reference_qa():
    for relative, expected_sheets in OPENVINO_WORKBOOKS.items():
        workbook = openpyxl.load_workbook(
            RELEASE_ROOT / relative,
            read_only=True,
            data_only=False,
            keep_links=True,
        )
        try:
            assert tuple(workbook.sheetnames) == expected_sheets
            known_sheets = set(workbook.sheetnames)
            formulas: list[str] = []
            literal_errors: list[tuple[str, str]] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        value = cell.value
                        if isinstance(value, str) and value.startswith("="):
                            formulas.append(value)
                        if isinstance(value, str) and value in FORMULA_ERROR_VALUES:
                            literal_errors.append((sheet.title, cell.coordinate))

            missing_sheet_refs = []
            external_refs = []
            for formula in formulas:
                if "[" in formula and "]" in formula:
                    external_refs.append(formula)
                for quoted, bare in re.findall(
                    r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_. ]*))!",
                    formula,
                ):
                    referenced = (quoted or bare).replace("''", "'")
                    if referenced not in known_sheets:
                        missing_sheet_refs.append((referenced, formula))

            assert not external_refs
            assert not missing_sheet_refs
            assert not literal_errors
            expected_formula_count = 275 if relative.startswith("04-") else 0
            assert len(formulas) == expected_formula_count
        finally:
            workbook.close()


def test_release_readiness_records_every_pdf_page_and_spreadsheet_qa():
    readiness = json.loads(
        (RELEASE_ROOT / "validation/release-readiness.json").read_text(
            encoding="utf-8"
        )
    )
    assert readiness["valid"] is True
    assert readiness["status"] == "ready_with_documented_limitations"
    assert readiness["qa_method"]["pdf_renderer"] == "PyMuPDF"

    pdf_qa = readiness["qa"]["pdf_reports"]
    expected_pdf_paths = {
        f"{route}/workbook/generated/{REPORT_STEMS[route]}.pdf" for route in ROUTES
    }
    assert set(pdf_qa) == expected_pdf_paths
    for relative, receipt in pdf_qa.items():
        with fitz.open(RELEASE_ROOT / relative) as document:
            assert receipt["page_count"] == document.page_count
            assert receipt["inspected_pages"] == list(range(1, document.page_count + 1))
            assert receipt["visual_review"] == "passed"
            assert all(page.get_text().strip() for page in document)

    workbook_qa = readiness["qa"]["openvino_workbooks"]
    assert set(workbook_qa) == set(OPENVINO_WORKBOOKS)
    for relative, receipt in workbook_qa.items():
        assert receipt["read_only"] is True
        assert tuple(receipt["sheets"]) == OPENVINO_WORKBOOKS[relative]
        assert receipt["missing_sheet_reference_count"] == 0
        assert receipt["external_formula_reference_count"] == 0
        assert receipt["formula_error_count"] == 0


def test_collection_validator_applies_blocking_release_metadata_gate():
    gate = validate_release_metadata(RELEASE_ROOT)
    report = validate_collection(RELEASE_ROOT)

    assert gate.valid is True
    assert report.gate("release_metadata").valid is True
    assert report.gate("release_readiness").valid is True
    assert report.valid is True


def test_release_metadata_gate_rejects_a_partial_metadata_publication(tmp_path):
    (tmp_path / "README.md").write_text("# Partial release\n", encoding="utf-8")

    gate = validate_release_metadata(tmp_path)

    assert gate.valid is False
    assert {issue.code for issue in gate.issues} == {"missing_release_metadata"}
