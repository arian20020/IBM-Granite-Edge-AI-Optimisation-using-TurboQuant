import sys
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from lxml import etree


sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from scripts.testing.final_results.markdown_renderer import render_markdown
from scripts.testing.final_results.report_model import (
    Report,
    ReportNote,
    ReportParagraph,
    ReportSection,
    ReportTable,
)


WORD_NAMESPACE = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _renderer():
    from scripts.testing.final_results.docx_renderer import render_docx

    return render_docx


def _parity_comparator():
    from scripts.testing.final_results.parity import compare_markdown_docx

    return compare_markdown_docx


def _report() -> Report:
    return Report(
        title="OpenVINO final results",
        route_id="openvino-official-upstream",
        revision="R2",
        generated_date=date(2026, 8, 30),
        evidence_ids=("EV-001", "EV-002"),
        sections=(
            ReportSection(
                title="Executive summary",
                blocks=(
                    ReportParagraph("Measurements are reported from normalized rows."),
                    ReportNote("Status colour is supplementary to explicit text."),
                    ReportTable(
                        table_id="T-01",
                        title="Execution status",
                        subtitle="Completed and unavailable configurations.",
                        columns=("Case", "Status", "Evidence"),
                        rows=(
                            ("granite|3b", "Passed", "EV-001"),
                            ("granite-8b", "Not collected", "EV-002<br>literal\nsee limitation"),
                            ("---", ":---:", "---"),
                        ),
                        footnotes=("Not collected means the campaign did not record the metric.",),
                    ),
                    ReportTable(
                        table_id="T-02",
                        title="Wide audit detail",
                        columns=("A", "B", "C", "D", "E", "F", "G"),
                        rows=(("1", "2", "3", "4", "5", "6", "Blocked"),),
                    ),
                    ReportParagraph("Narrative resumes in portrait orientation."),
                ),
            ),
        ),
    )


def _word_xml(docx_path: Path, member: str) -> etree._Element:
    with zipfile.ZipFile(docx_path) as package:
        return etree.fromstring(package.read(member))


def test_render_docx_applies_professional_accessible_document_structure(tmp_path):
    output = tmp_path / "report.docx"

    _renderer()(_report(), output)

    document_xml = _word_xml(output, "word/document.xml")
    fills = document_xml.xpath("//w:shd/@w:fill", namespaces=WORD_NAMESPACE)
    assert "0F2747" in fills
    assert "0F766E" in fills
    assert "0B63CE" in fills

    document = Document(output)
    table_text = [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]
    assert ["granite|3b", "Passed", "EV-001"] in table_text[1]
    assert ["granite-8b", "Not collected", "EV-002<br>literal\nsee limitation"] in table_text[1]
    assert ["---", ":---:", "---"] in table_text[1]
    assert ["1", "2", "3", "4", "5", "6", "Blocked"] in table_text[2]

    rows = document_xml.xpath("//w:tbl/w:tr", namespaces=WORD_NAMESPACE)
    assert rows
    assert all(row.xpath("./w:trPr/w:cantSplit", namespaces=WORD_NAMESPACE) for row in rows)
    table_headers = document_xml.xpath("//w:tbl/w:tr[1]/w:trPr/w:tblHeader", namespaces=WORD_NAMESPACE)
    assert len(table_headers) == len(document.tables)

    footer_parts = []
    with zipfile.ZipFile(output) as package:
        for name in package.namelist():
            if name.startswith("word/footer") and name.endswith(".xml"):
                footer_parts.append(package.read(name).decode("utf-8"))
    combined_footer = "".join(footer_parts)
    assert "openvino-official-upstream" in combined_footer
    assert "R2" in combined_footer
    assert "PAGE" in combined_footer


def test_render_docx_uses_landscape_only_beyond_the_wide_table_threshold(tmp_path):
    output = tmp_path / "report.docx"

    _renderer()(_report(), output)

    document = Document(output)
    orientations = [section.orientation for section in document.sections]
    assert orientations == [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT]

    def report_with_estimated_width(widths: tuple[int, ...]) -> Report:
        return Report(
            title="Width boundary",
            route_id="width-test",
            revision="R1",
            generated_date=date(2026, 8, 30),
            sections=(
                ReportSection(
                    title="Results",
                    blocks=(
                        ReportTable(
                            table_id="T-01",
                            title="Boundary table",
                            columns=tuple("X" * width for width in widths),
                            rows=(tuple("1" for _ in widths),),
                        ),
                        ReportParagraph("Orientation reset marker."),
                    ),
                ),
            ),
        )

    at_boundary = tmp_path / "at-boundary.docx"
    beyond_boundary = tmp_path / "beyond-boundary.docx"
    _renderer()(report_with_estimated_width((15, 16, 16, 16, 16, 16)), at_boundary)
    _renderer()(report_with_estimated_width((16, 16, 16, 16, 16, 16)), beyond_boundary)

    assert [section.orientation for section in Document(at_boundary).sections] == [
        WD_ORIENT.PORTRAIT
    ]
    assert [section.orientation for section in Document(beyond_boundary).sections] == [
        WD_ORIENT.PORTRAIT,
        WD_ORIENT.LANDSCAPE,
        WD_ORIENT.PORTRAIT,
    ]


def test_render_docx_styles_the_controlled_artifact_unavailable_label_neutral_grey(tmp_path):
    report = Report(
        title="Unavailable status",
        route_id="status-test",
        revision="R1",
        generated_date=date(2026, 8, 30),
        sections=(
            ReportSection(
                title="Availability",
                blocks=(
                    ReportTable(
                        table_id="T-01",
                        title="Statuses",
                        columns=("Case", "Status"),
                        rows=(("granite-30b", "Artifact unavailable"),),
                    ),
                ),
            ),
        ),
    )
    output = tmp_path / "artifact-unavailable.docx"

    _renderer()(report, output)

    document = Document(output)
    status_cell = document.tables[1].cell(1, 1)
    shading = status_cell._tc.get_or_add_tcPr().find("{%s}shd" % WORD_NAMESPACE["w"])
    assert status_cell.text == "Artifact unavailable"
    assert shading is not None
    assert shading.get("{%s}fill" % WORD_NAMESPACE["w"]) == "E7E6E6"


def test_render_docx_keeps_status_headers_blue_and_styles_only_data_status_cells(tmp_path):
    statuses = ("Passed", "Failed", "Blocked", "Artifact unavailable")
    report = Report(
        title="Status contrast",
        route_id="status-test",
        revision="R1",
        generated_date=date(2026, 8, 30),
        sections=(
            ReportSection(
                title="Status table",
                blocks=(
                    ReportTable(
                        table_id="T-01",
                        title="Status header and data distinction",
                        columns=statuses,
                        rows=(statuses,),
                    ),
                ),
            ),
        ),
    )
    output = tmp_path / "status-header-and-data.docx"

    _renderer()(report, output)

    table = Document(output).tables[1]
    expected_data_fills = ("E2F0D9", "FCE4D6", "FFF2CC", "E7E6E6")

    def fill(cell) -> str:
        shading = cell._tc.get_or_add_tcPr().find("{%s}shd" % WORD_NAMESPACE["w"])
        assert shading is not None
        return shading.get("{%s}fill" % WORD_NAMESPACE["w"])

    def colours(cell) -> set[str]:
        return {
            str(run.font.color.rgb)
            for paragraph in cell.paragraphs
            for run in paragraph.runs
        }

    def contrast_ratio(first: str, second: str) -> float:
        def luminance(colour: str) -> float:
            channels = [int(colour[index : index + 2], 16) / 255 for index in (0, 2, 4)]
            linear = [
                channel / 12.92
                if channel <= 0.04045
                else ((channel + 0.055) / 1.055) ** 2.4
                for channel in channels
            ]
            return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]

        light, dark = sorted((luminance(first), luminance(second)), reverse=True)
        return (light + 0.05) / (dark + 0.05)

    for cell in table.rows[0].cells:
        assert fill(cell) == "0B63CE"
        assert colours(cell) == {"FFFFFF"}
        assert contrast_ratio("0B63CE", "FFFFFF") >= 4.5

    for cell, expected_fill in zip(table.rows[1].cells, expected_data_fills, strict=True):
        assert fill(cell) == expected_fill
        assert colours(cell) == {"1F1F1F"}
        assert contrast_ratio(expected_fill, "1F1F1F") >= 4.5


def test_render_docx_normalizes_package_timestamps_and_bytes(tmp_path):
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"

    _renderer()(_report(), first)
    _renderer()(_report(), second)

    with zipfile.ZipFile(first) as package:
        assert package.infolist()
        assert {info.date_time for info in package.infolist()} == {(1980, 1, 1, 0, 0, 0)}
    assert first.read_bytes() == second.read_bytes()


def test_compare_markdown_docx_reports_exact_heading_and_table_cell_parity(tmp_path):
    markdown = tmp_path / "report.md"
    docx = tmp_path / "report.docx"
    render_markdown(_report(), markdown)
    _renderer()(_report(), docx)

    result = _parity_comparator()(markdown, docx)

    assert result["matches"] is True
    assert result["headings_match"] is True
    assert result["tables_match"] is True
    assert result["headings"]["markdown"] == [
        "OpenVINO final results",
        "Executive summary",
        "T-01 — Execution status",
        "T-02 — Wide audit detail",
    ]
    assert result["headings"]["markdown"] == result["headings"]["docx"]
    assert result["tables"]["markdown"] == result["tables"]["docx"]
    assert result["differences"] == []


def test_compare_markdown_docx_identifies_a_changed_cell(tmp_path):
    markdown = tmp_path / "report.md"
    docx = tmp_path / "report.docx"
    render_markdown(_report(), markdown)
    _renderer()(_report(), docx)
    changed = Document(docx)
    changed.tables[1].cell(1, 1).text = "Failed"
    changed.save(docx)

    result = _parity_comparator()(markdown, docx)

    assert result["matches"] is False
    assert result["headings_match"] is True
    assert result["tables_match"] is False
    assert result["differences"] == [
        {
            "kind": "tables",
            "markdown": result["tables"]["markdown"],
            "docx": result["tables"]["docx"],
        }
    ]


def test_compare_markdown_docx_preserves_literal_backslashes_in_cells(tmp_path):
    report = Report(
        title="Path report",
        route_id="path-test",
        revision="R1",
        generated_date=date(2026, 8, 30),
        sections=(
            ReportSection(
                title="Evidence",
                blocks=(
                    ReportTable(
                        table_id="T-01",
                        title="Paths",
                        columns=("Evidence path",),
                        rows=((r"evidence\run-001\result.json",),),
                    ),
                ),
            ),
        ),
    )
    markdown = tmp_path / "report.md"
    docx = tmp_path / "report.docx"
    render_markdown(report, markdown)
    _renderer()(report, docx)

    result = _parity_comparator()(markdown, docx)

    assert result["matches"] is True
