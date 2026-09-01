from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl


sys.path.insert(0, str(Path(__file__).resolve().parents[4]))


from scripts.testing.reporting.workbook_portability import (
    sanitize_workbook,
    write_portable_openvino_route_workbook,
    write_portable_openvino_workbooks,
)


WINDOWS_ABSOLUTE = re.compile(r"^[A-Za-z]:[\\/]")

REPOSITORY_ROOT = Path(__file__).resolve().parents[4]
WORKBOOK_CASES = (
    {
        "source": (
            "04-openvino-experimental-fork/results/source/"
            "Granite_OpenVINO_Final_Healthcare_Education_Results_2026-08-30.xlsx"
        ),
        "portable": (
            "04-openvino-experimental-fork/workbook/generated/"
            "openvino-experimental-fork-portable-results.xlsx"
        ),
        "receipt": (
            "04-openvino-experimental-fork/workbook/generated/"
            "portable-workbook-provenance.json"
        ),
        "sheets": (
            "Dashboard",
            "Format Comparison",
            "Sector Summary",
            "Detailed Results",
            "Quality Details",
            "Availability Matrix",
            "Methodology",
            "Source Data",
        ),
        "dimensions": {
            "Dashboard": "A1:N28",
            "Format Comparison": "A1:M17",
            "Sector Summary": "A1:N32",
            "Detailed Results": "A1:AF82",
            "Quality Details": "A1:U3889",
            "Availability Matrix": "A1:J20",
            "Methodology": "A1:H37",
            "Source Data": "A1:BB83",
        },
        "replacement_count": 55,
        "formula_count": 275,
    },
    {
        "source": (
            "05-openvino-official-upstream/results/source/"
            "Granite_Official_OpenVINO_TurboQuant_Results_2026-08-30_v2_Missing_Attempts.xlsx"
        ),
        "portable": (
            "05-openvino-official-upstream/workbook/generated/"
            "openvino-official-upstream-portable-results.xlsx"
        ),
        "receipt": (
            "05-openvino-official-upstream/workbook/generated/"
            "portable-workbook-provenance.json"
        ),
        "sheets": (
            "Executive Summary",
            "Format Comparison",
            "Sector Summary",
            "Detailed Results",
            "Quality Details",
            "Availability Matrix",
            "Missing Attempts",
            "Methodology",
            "Source Data",
        ),
        "dimensions": {
            "Executive Summary": "A1:L29",
            "Format Comparison": "A1:M14",
            "Sector Summary": "A1:I185",
            "Detailed Results": "A1:AJ50",
            "Quality Details": "A1:U2165",
            "Availability Matrix": "A1:G14",
            "Missing Attempts": "A1:K11",
            "Methodology": "A1:H31",
            "Source Data": "A1:E16",
        },
        "replacement_count": 15,
        "formula_count": 0,
    },
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_sanitizer_replaces_machine_paths_without_changing_workbook_semantics(
    tmp_path: Path,
):
    source = tmp_path / "source.xlsx"
    first_output = tmp_path / "portable-1.xlsx"
    second_output = tmp_path / "portable-2.xlsx"

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet["A1"] = (
        r"C:\Users\Student\IBM-Granite-Edge-AI-Optimisation-using-TurboQuant"
        r"\.worktrees\openvino-turboquant-recovery\experiments\raw-results\run.json"
    )
    sheet["B1"] = "=SUM(1,2)"
    sheet["C1"] = 12.5
    sheet["D1"] = "Evidence"
    sheet["D1"].hyperlink = (
        r"C:\Users\Student\IBM-Granite-Edge-AI-Optimisation-using-TurboQuant"
        r"\.worktrees\openvino-turboquant-recovery\docs\testing\final-results\README.md"
    )
    sheet.merge_cells("E1:F1")
    sheet["E1"] = "Merged"
    workbook.create_sheet("Methods")["A1"] = "unchanged"
    workbook.save(source)
    source_hash = _sha256(source)

    first_receipt = sanitize_workbook(source, first_output)
    second_receipt = sanitize_workbook(source, second_output)

    assert _sha256(source) == source_hash
    assert _sha256(first_output) == _sha256(second_output)
    assert first_receipt["source_sha256"] == source_hash
    assert first_receipt["output_sha256"] == second_receipt["output_sha256"]
    assert first_receipt["replacement_count"] == 2
    assert first_receipt["formula_count"] == 1

    portable = openpyxl.load_workbook(first_output, data_only=False, keep_links=True)
    try:
        assert portable.sheetnames == ["Results", "Methods"]
        assert portable["Results"]["A1"].value == "experiments/raw-results/run.json"
        assert portable["Results"]["B1"].value == "=SUM(1,2)"
        assert portable["Results"]["C1"].value == 12.5
        assert portable["Results"]["D1"].hyperlink.target == (
            "docs/testing/final-results/README.md"
        )
        assert {str(item) for item in portable["Results"].merged_cells.ranges} == {
            "E1:F1"
        }
        for worksheet in portable.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        assert not WINDOWS_ABSOLUTE.match(cell.value)
                    if cell.hyperlink and isinstance(cell.hyperlink.target, str):
                        assert not WINDOWS_ABSOLUTE.match(cell.hyperlink.target)
    finally:
        portable.close()


def test_real_openvino_portable_workbooks_preserve_values_and_publish_provenance(
    tmp_path: Path,
):
    collection = tmp_path / "docs/testing/final-results"
    source_hashes: dict[str, str] = {}
    for case in WORKBOOK_CASES:
        relative = str(case["source"])
        source = REPOSITORY_ROOT / "docs/testing/final-results" / relative
        target = collection / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
        source_hashes[relative] = _sha256(target)

    receipts = write_portable_openvino_workbooks(tmp_path)

    assert len(receipts) == 2
    for case, returned in zip(WORKBOOK_CASES, receipts, strict=True):
        source_relative = str(case["source"])
        portable_relative = str(case["portable"])
        receipt_relative = str(case["receipt"])
        source_path = collection / source_relative
        portable_path = collection / portable_relative
        receipt_path = collection / receipt_relative

        assert _sha256(source_path) == source_hashes[source_relative]
        assert returned == json.loads(receipt_path.read_text(encoding="utf-8"))
        assert returned["source_path"] == (
            f"docs/testing/final-results/{source_relative}"
        )
        assert returned["output_path"] == (
            f"docs/testing/final-results/{portable_relative}"
        )
        assert returned["receipt_path"] == (
            f"docs/testing/final-results/{receipt_relative}"
        )
        assert returned["replacement_count"] == case["replacement_count"]
        assert returned["formula_count"] == case["formula_count"]
        assert returned["cached_formula_value_count"] == 0
        assert returned["source_sha256"] == source_hashes[source_relative]
        assert returned["output_sha256"] == _sha256(portable_path)

        source_book = openpyxl.load_workbook(
            source_path, data_only=False, keep_links=True
        )
        portable_book = openpyxl.load_workbook(
            portable_path, data_only=False, keep_links=True
        )
        try:
            assert tuple(portable_book.sheetnames) == case["sheets"]
            assert {
                sheet.title: sheet.calculate_dimension()
                for sheet in portable_book.worksheets
            } == case["dimensions"]
            formulas = 0
            machine_paths = []
            for source_sheet, portable_sheet in zip(
                source_book.worksheets, portable_book.worksheets, strict=True
            ):
                for row in source_sheet.iter_rows():
                    for source_cell in row:
                        portable_cell = portable_sheet[source_cell.coordinate]
                        if isinstance(source_cell.value, str) and WINDOWS_ABSOLUTE.match(
                            source_cell.value
                        ):
                            assert portable_cell.value.startswith(
                                ("experiments/", "external/")
                            )
                            assert "\\" not in portable_cell.value
                        else:
                            assert portable_cell.value == source_cell.value
                        if isinstance(portable_cell.value, str):
                            formulas += portable_cell.value.startswith("=")
                            if WINDOWS_ABSOLUTE.match(portable_cell.value):
                                machine_paths.append(
                                    f"{portable_sheet.title}!{portable_cell.coordinate}"
                                )
                        if portable_cell.hyperlink and WINDOWS_ABSOLUTE.match(
                            portable_cell.hyperlink.target
                        ):
                            machine_paths.append(
                                f"{portable_sheet.title}!{portable_cell.coordinate}:hyperlink"
                            )
            assert formulas == case["formula_count"]
            assert machine_paths == []
        finally:
            source_book.close()
            portable_book.close()


def test_route_writer_publishes_only_the_selected_openvino_derivative(tmp_path: Path):
    route = tmp_path / "docs/testing/final-results/04-openvino-experimental-fork"
    case = WORKBOOK_CASES[0]
    source_relative = Path(str(case["source"])).relative_to(route.name)
    source = REPOSITORY_ROOT / "docs/testing/final-results" / str(case["source"])
    target = route / source_relative
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, target)

    receipt = write_portable_openvino_route_workbook(tmp_path, route)

    assert receipt["replacement_count"] == 55
    assert (route / Path(str(case["portable"])).relative_to(route.name)).is_file()
    assert (route / Path(str(case["receipt"])).relative_to(route.name)).is_file()
    assert not (
        tmp_path
        / "docs/testing/final-results/05-openvino-official-upstream/workbook/generated/"
        "openvino-official-upstream-portable-results.xlsx"
    ).exists()
