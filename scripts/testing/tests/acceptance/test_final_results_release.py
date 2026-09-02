from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import subprocess
import sys
import tarfile
from pathlib import Path, PurePosixPath
from urllib.parse import urlsplit

import fitz
import openpyxl
import pytest


sys.path.insert(0, str(Path(__file__).resolve().parents[4]))


from scripts.testing.reporting.validate import (
    validate_collection,
    validate_release_metadata,
)


REPOSITORY_ROOT = Path(__file__).resolve().parents[4]
RELEASE_ROOT = REPOSITORY_ROOT / "docs/testing/final-results"
PRECOMPACT_LAYOUT_COMMIT = "71fec68a12308d89ae90fc80f17302bd3527845c"
ROUTES = (
    "01-upstream-llama-cpp",
    "02-atomicbot-turboquant",
    "03-animehacker-tq3-0",
    "04-openvino-experimental-fork",
    "05-openvino-official-upstream",
    "06-cross-route-comparison",
)
REPORT_STEMS = {
    "01-upstream-llama-cpp": "upstream-llama-cpp-report",
    "02-atomicbot-turboquant": "atomicbot-turboquant-report",
    "03-animehacker-tq3-0": "animehacker-tq3-0-report",
    "04-openvino-experimental-fork": "openvino-experimental-fork-report",
    "05-openvino-official-upstream": "openvino-official-upstream-report",
    "06-cross-route-comparison": "cross-route-comparison-report",
}
OPENVINO_SOURCE_WORKBOOKS = {
    "04-openvino-experimental-fork/evidence/source/Granite_OpenVINO_Final_Healthcare_Education_Results_2026-08-30.xlsx": (
        "Dashboard",
        "Format Comparison",
        "Sector Summary",
        "Detailed Results",
        "Quality Details",
        "Availability Matrix",
        "Methodology",
        "Source Data",
    ),
    "05-openvino-official-upstream/evidence/source/Granite_Official_OpenVINO_TurboQuant_Results_2026-08-30_v2_Missing_Attempts.xlsx": (
        "Executive Summary",
        "Format Comparison",
        "Sector Summary",
        "Detailed Results",
        "Quality Details",
        "Availability Matrix",
        "Missing Attempts",
        "Methodology",
        "Source Data",
    ),
}
OPENVINO_PORTABLE_WORKBOOKS = {
    "04-openvino-experimental-fork/reports/openvino-experimental-fork-results.xlsx": OPENVINO_SOURCE_WORKBOOKS[
        "04-openvino-experimental-fork/evidence/source/Granite_OpenVINO_Final_Healthcare_Education_Results_2026-08-30.xlsx"
    ],
    "05-openvino-official-upstream/reports/openvino-official-upstream-results.xlsx": OPENVINO_SOURCE_WORKBOOKS[
        "05-openvino-official-upstream/evidence/source/Granite_Official_OpenVINO_TurboQuant_Results_2026-08-30_v2_Missing_Attempts.xlsx"
    ],
}
OPENVINO_WORKBOOKS = {
    **OPENVINO_SOURCE_WORKBOOKS,
    **OPENVINO_PORTABLE_WORKBOOKS,
}
OPENVINO_WORKBOOK_RECEIPTS = {
    "04-openvino-experimental-fork/reports/openvino-experimental-fork-results-provenance.json",
    "05-openvino-official-upstream/reports/openvino-official-upstream-results-provenance.json",
}
OPENVINO_MACHINE_PATH_COUNTS = {
    **{path: count for path, count in zip(OPENVINO_SOURCE_WORKBOOKS, (55, 15), strict=True)},
    **{path: 0 for path in OPENVINO_PORTABLE_WORKBOOKS},
}
CATALOGS = (
    "catalog/route-register.csv",
    "catalog/campaign-summary.csv",
    "catalog/performance-summary.csv",
    "catalog/quality-summary.csv",
    "catalog/failure-summary.csv",
    "catalog/evidence-manifest.csv",
    "catalog/claim-evidence-map.csv",
    "catalog/comparability-matrix.csv",
)
FORMULA_ERROR_VALUES = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
}


def _git_index_release_blobs() -> dict[str, bytes]:
    prefix = RELEASE_ROOT.relative_to(REPOSITORY_ROOT).as_posix()
    listing = subprocess.run(
        ["git", "ls-files", "--stage", "--", prefix],
        cwd=REPOSITORY_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    blobs: dict[str, bytes] = {}
    for line in listing.splitlines():
        metadata, repository_relative = line.split("\t", 1)
        _, object_id, stage = metadata.split()
        assert stage == "0", repository_relative
        release_relative = PurePosixPath(repository_relative).relative_to(prefix).as_posix()
        blobs[release_relative] = subprocess.run(
            ["git", "cat-file", "blob", object_id],
            cwd=REPOSITORY_ROOT,
            check=True,
            capture_output=True,
        ).stdout
    return blobs


def _published_evidence_sources() -> dict[str, tuple[str, int]]:
    records: dict[str, tuple[str, int]] = {}
    for route in ROUTES[:5]:
        index = RELEASE_ROOT / route / "evidence/evidence-index.csv"
        with index.open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                relative = row["relative_path"].replace("\\", "/")
                identity = (row["sha256"].lower(), int(row["size_bytes"]))
                assert records.get(relative, identity) == identity, relative
                records[relative] = identity
    return records


def _write_filesystem_manifest(release_root: Path) -> None:
    manifest = release_root / "manifest-sha256.txt"
    entries = []
    for path in sorted(
        (path for path in release_root.rglob("*") if path.is_file() and path != manifest),
        key=lambda path: path.relative_to(release_root).as_posix(),
    ):
        relative = path.relative_to(release_root).as_posix()
        entries.append(f"{hashlib.sha256(path.read_bytes()).hexdigest()}  {relative}")
    manifest.write_text("\n".join(entries) + "\n", encoding="utf-8", newline="\n")


@pytest.fixture
def canonical_release(tmp_path: Path) -> Path:
    root = tmp_path / "final-results"
    for relative, payload in _git_index_release_blobs().items():
        path = root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload)
    _write_filesystem_manifest(root)
    baseline = validate_release_metadata(root)
    assert baseline.valid, baseline.issues
    return root


def _markdown_targets(path: Path) -> set[str]:
    text = path.read_text(encoding="utf-8")
    return {
        target.split("#", 1)[0]
        for target in re.findall(r"\[[^\]]+\]\(([^)]+)\)", text)
        if target and not urlsplit(target).scheme
    }


def _type_names(entity: dict[str, object]) -> set[str]:
    value = entity.get("@type", ())
    if isinstance(value, str):
        return {value}
    return {str(item) for item in value}


def _graph() -> tuple[dict[str, object], dict[str, dict[str, object]]]:
    crate = json.loads(
        (RELEASE_ROOT / "ro-crate-metadata.json").read_text(encoding="utf-8")
    )
    entities = {str(entity["@id"]): entity for entity in crate["@graph"]}
    return crate, entities


def _path_migration_rows() -> list[dict[str, str]]:
    path = REPOSITORY_ROOT / "docs/testing/cleanup/PATH-MIGRATION.csv"
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _removed_precompact_published_paths() -> set[str]:
    prefix = RELEASE_ROOT.relative_to(REPOSITORY_ROOT).as_posix() + "/"
    before = subprocess.run(
        [
            "git",
            "ls-tree",
            "-r",
            "--name-only",
            PRECOMPACT_LAYOUT_COMMIT,
            "--",
            prefix,
        ],
        cwd=REPOSITORY_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.splitlines()
    current = subprocess.run(
        ["git", "ls-files", "--", prefix],
        cwd=REPOSITORY_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.splitlines()
    route_prefixes = tuple(f"{prefix}{route}/" for route in ROUTES)
    return {
        path
        for path in set(before) - set(current)
        if path.startswith(route_prefixes)
    }


def test_release_portal_links_every_route_report_workbook_and_control_surface():
    targets = _markdown_targets(RELEASE_ROOT / "README.md")
    expected = {
        "CHANGELOG.md",
        "REPRODUCING.md",
        "LICENSES.md",
        "ro-crate-metadata.json",
        "manifest-sha256.txt",
        "validation/release-readiness.json",
        *CATALOGS,
    }
    for route in ROUTES:
        stem = REPORT_STEMS[route]
        expected.update(
            {
                f"{route}/README.md",
                f"{route}/data/route.json",
                f"{route}/evidence/claim-evidence-map.csv",
                f"{route}/validation/validation.md",
                f"{route}/reproduction/README.md",
                f"{route}/reports/{stem}.md",
                f"{route}/reports/{stem}.docx",
                f"{route}/reports/{stem}.pdf",
            }
        )
    expected.update(OPENVINO_WORKBOOKS)
    expected.update(OPENVINO_WORKBOOK_RECEIPTS)

    assert expected <= targets
    assert all((RELEASE_ROOT / target).is_file() for target in expected)


def test_ro_crate_is_version_13_and_all_packaged_data_entity_ids_are_relative():
    crate, entities = _graph()

    assert crate["@context"][0] == "https://w3id.org/ro/crate/1.3/context"
    descriptor = entities["ro-crate-metadata.json"]
    assert descriptor["conformsTo"] == {"@id": "https://w3id.org/ro/crate/1.3"}
    assert descriptor["about"] == {"@id": "./"}
    assert "Dataset" in _type_names(entities["./"])

    for entity_id, entity in entities.items():
        if not (_type_names(entity) & {"File", "Dataset"}):
            continue
        parsed = urlsplit(entity_id)
        assert not parsed.scheme, entity_id
        assert not parsed.netloc, entity_id
        assert not entity_id.startswith(("/", "\\")), entity_id
        assert not re.match(r"^[A-Za-z]:", entity_id), entity_id
        assert ".." not in PurePosixPath(entity_id).parts, entity_id
        if "File" in _type_names(entity):
            assert (RELEASE_ROOT / entity_id).is_file(), entity_id

    old_published = {
        row["old_path"].removeprefix("docs/testing/final-results/")
        for row in _path_migration_rows()
        if row["old_path"].startswith("docs/testing/final-results/")
    }
    assert not (old_published & entities.keys())


def test_ro_crate_covers_reports_tables_systems_and_generation_provenance():
    _, entities = _graph()
    expected_files = set(CATALOGS)
    for route in ROUTES:
        stem = REPORT_STEMS[route]
        expected_files.update(
            {
                f"{route}/reports/{stem}.md",
                f"{route}/reports/{stem}.docx",
                f"{route}/reports/{stem}.pdf",
            }
        )
    for route in ROUTES[:5]:
        expected_files.update(
            {
                f"{route}/data/attempts.csv",
                f"{route}/data/measurements.csv",
                f"{route}/data/summaries.csv",
            }
        )
    expected_files.update(OPENVINO_WORKBOOKS)
    expected_files.update(OPENVINO_WORKBOOK_RECEIPTS)
    assert expected_files <= entities.keys()

    for route in ROUTES[:5]:
        assert f"#{route}-repository" in entities
        assert f"#{route}-hardware" in entities
        assert f"#{route}-software" in entities

    generated = {
        entity_id: entity
        for entity_id, entity in entities.items()
        if "#generated-artifact" in {
            str(value.get("@id"))
            for value in entity.get("additionalType", ())
            if isinstance(value, dict)
        }
    }
    assert {
        f"{route}/reports/{REPORT_STEMS[route]}.{suffix}"
        for route in ROUTES
        for suffix in ("docx", "pdf")
    } <= generated.keys()
    assert "manifest-sha256.txt" in generated
    assert set(OPENVINO_PORTABLE_WORKBOOKS) <= generated.keys()
    assert OPENVINO_WORKBOOK_RECEIPTS <= generated.keys()

    for entity_id, entity in generated.items():
        activity_id = entity["wasGeneratedBy"]["@id"]
        activity = entities[activity_id]
        assert "CreateAction" in _type_names(activity)
        assert {item["@id"] for item in activity["result"]} >= {entity_id}
        assert activity.get("object"), entity_id


def test_release_manifest_is_complete_sorted_self_excluding_and_hash_valid():
    manifest = RELEASE_ROOT / "manifest-sha256.txt"
    lines = manifest.read_text(encoding="utf-8").splitlines()
    entries: dict[str, str] = {}
    for line in lines:
        match = re.fullmatch(r"([0-9a-f]{64})  ([^\\]+)", line)
        assert match, line
        digest, relative = match.groups()
        assert relative not in entries
        entries[relative] = digest

    expected = {
        path.relative_to(RELEASE_ROOT).as_posix()
        for path in RELEASE_ROOT.rglob("*")
        if path.is_file() and path != manifest
    }
    assert list(entries) == sorted(entries)
    assert "manifest-sha256.txt" not in entries
    assert set(entries) == expected
    for relative, expected_digest in entries.items():
        actual = hashlib.sha256((RELEASE_ROOT / relative).read_bytes()).hexdigest()
        assert actual == expected_digest, relative


def test_licensing_gaps_are_explicit_and_citation_is_not_invented():
    text = (RELEASE_ROOT / "LICENSES.md").read_text(encoding="utf-8").casefold()

    assert "no repository-level license file" in text
    assert "no license grant" in text
    assert "not verified" in text
    assert all(route in text for route in ROUTES[:5])
    assert not (RELEASE_ROOT / "CITATION.cff").exists()
    assert "citation.cff" in text
    assert "verified author metadata" in text


def test_reproduction_commands_are_exact_and_preserve_the_no_rerun_boundary():
    text = (RELEASE_ROOT / "REPRODUCING.md").read_text(encoding="utf-8")

    assert "--validate-only" in text
    assert "test_final_results_*.py" in text
    assert "does not rerun benchmarks" in text.casefold()
    assert "does not modify raw evidence" in text.casefold()
    assert "no external hydration" in text.casefold()
    assert "non-calculating" in text.casefold()
    assert "recalculation" in text.casefold()
    for route in ROUTES:
        assert f"{route}/reproduction/" in text

    reproduction_docs = [RELEASE_ROOT / "REPRODUCING.md"] + sorted(
        RELEASE_ROOT.glob("*/reproduction/**/*.md")
    )
    command_lines = []
    for path in reproduction_docs:
        for block in re.findall(
            r"```powershell\s*\n(.*?)```",
            path.read_text(encoding="utf-8"),
            re.DOTALL | re.IGNORECASE,
        ):
            command_lines.extend(
                line.strip()
                for line in block.splitlines()
                if line.strip()
                and not line.lstrip().startswith(("#", "$releaseTests"))
                and not line.lstrip().startswith(("(", ")", "+"))
            )

    assert command_lines
    assert not any("scripts.testing.reporting" in line for line in command_lines)
    assert not any("scripts/testing/build_final_results.py" in line for line in command_lines)
    for line in command_lines:
        if re.search(r"(?:python(?:\.exe)?['\"]?\s+)(?!-m pytest)", line, re.IGNORECASE):
            assert "-m scripts.testing.cli." in line, line


def test_path_migration_covers_every_removed_precompact_published_path():
    rows = _path_migration_rows()
    mappings = {row["old_path"]: row["new_path"] for row in rows}
    removed = _removed_precompact_published_paths()

    assert removed
    assert removed <= mappings.keys()
    for old_path in sorted(removed):
        assert not (REPOSITORY_ROOT / old_path).exists(), old_path
        assert (REPOSITORY_ROOT / mappings[old_path]).is_file(), mappings[old_path]


def test_openvino_workbooks_pass_read_only_sheet_and_formula_reference_qa():
    for relative, expected_sheets in OPENVINO_WORKBOOKS.items():
        workbook = openpyxl.load_workbook(
            RELEASE_ROOT / relative,
            read_only=False,
            data_only=False,
            keep_links=True,
        )
        try:
            assert tuple(workbook.sheetnames) == expected_sheets
            known_sheets = set(workbook.sheetnames)
            formulas: list[str] = []
            literal_errors: list[tuple[str, str]] = []
            machine_paths: list[tuple[str, str]] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        value = cell.value
                        if isinstance(value, str) and value.startswith("="):
                            formulas.append(value)
                        if isinstance(value, str) and value in FORMULA_ERROR_VALUES:
                            literal_errors.append((sheet.title, cell.coordinate))
                        if isinstance(value, str) and (
                            re.match(r"^[A-Za-z]:[\\/]", value)
                            or value.casefold().startswith("file://")
                            or value.startswith("\\\\")
                        ):
                            machine_paths.append((sheet.title, cell.coordinate))
                        if cell.hyperlink and isinstance(cell.hyperlink.target, str) and (
                            re.match(r"^[A-Za-z]:[\\/]", cell.hyperlink.target)
                            or cell.hyperlink.target.casefold().startswith("file://")
                            or cell.hyperlink.target.startswith("\\\\")
                        ):
                            machine_paths.append((sheet.title, f"{cell.coordinate}:hyperlink"))

            missing_sheet_refs = []
            external_refs = []
            for formula in formulas:
                if "[" in formula and "]" in formula:
                    external_refs.append(formula)
                for quoted, bare in re.findall(
                    r"(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_. ]*))!",
                    formula,
                ):
                    referenced = (quoted or bare).replace("''", "'")
                    if referenced not in known_sheets:
                        missing_sheet_refs.append((referenced, formula))

            assert not external_refs
            assert not missing_sheet_refs
            assert not literal_errors
            expected_formula_count = 275 if relative.startswith("04-") else 0
            assert len(formulas) == expected_formula_count
            assert len(machine_paths) == OPENVINO_MACHINE_PATH_COUNTS[relative]
        finally:
            workbook.close()


def test_release_readiness_records_every_pdf_page_and_spreadsheet_qa():
    readiness = json.loads(
        (RELEASE_ROOT / "validation/release-readiness.json").read_text(
            encoding="utf-8"
        )
    )
    assert readiness["valid"] is True
    assert readiness["status"] == "ready_with_documented_limitations"
    assert readiness["release_version"] == "unified-final-results-2026-09-01-v2"
    assert readiness["basis"]["raw_git_archive_self_contained"] is True
    assert readiness["basis"]["external_hydration_required"] is False
    assert readiness["basis"]["published_evidence_source_count"] == 1846
    assert readiness["basis"]["published_evidence_relationship_record_count"] == 1857
    assert readiness["basis"]["semantic_snapshot_relationship_count"] == 1846
    assert readiness["basis"]["published_stale_migrated_path_count"] == 0
    assert readiness["basis"]["manifest_generation_order"].endswith(
        "self-excluding staged-blob manifest last"
    )
    assert readiness["qa_method"]["pdf_renderer"] == "PyMuPDF"

    pdf_qa = readiness["qa"]["pdf_reports"]
    expected_pdf_paths = {
        f"{route}/reports/{REPORT_STEMS[route]}.pdf" for route in ROUTES
    }
    assert set(pdf_qa) == expected_pdf_paths
    for relative, receipt in pdf_qa.items():
        with fitz.open(RELEASE_ROOT / relative) as document:
            assert receipt["page_count"] == document.page_count
            assert receipt["inspected_pages"] == list(range(1, document.page_count + 1))
            assert receipt["visual_review"] == "passed"
            assert all(page.get_text().strip() for page in document)

    workbook_qa = readiness["qa"]["openvino_workbooks"]
    assert set(workbook_qa) == set(OPENVINO_WORKBOOKS)
    for relative, receipt in workbook_qa.items():
        assert receipt["read_only"] is True
        assert tuple(receipt["sheets"]) == OPENVINO_WORKBOOKS[relative]
        assert receipt["missing_sheet_reference_count"] == 0
        assert receipt["external_formula_reference_count"] == 0
        assert receipt["formula_error_count"] == 0
        assert receipt["machine_absolute_path_count"] == OPENVINO_MACHINE_PATH_COUNTS[relative]
        assert receipt["cached_formula_value_count"] == 0
        expected_role = (
            "primary_portable_derivative"
            if relative in OPENVINO_PORTABLE_WORKBOOKS
            else "immutable_evidence_only_nonportable"
        )
        assert receipt["role"] == expected_role

    limitations = " ".join(readiness["limitations"]).casefold()
    assert "non-calculating" in limitations
    assert "recalculation" in limitations


def test_collection_validator_applies_blocking_release_metadata_gate():
    gate = validate_release_metadata(RELEASE_ROOT)
    report = validate_collection(RELEASE_ROOT)

    assert gate.valid is True
    assert report.gate("release_metadata").valid is True
    assert report.gate("release_readiness").valid is True
    assert report.valid is True


def test_release_metadata_gate_rejects_a_partial_metadata_publication(tmp_path):
    (tmp_path / "README.md").write_text("# Partial release\n", encoding="utf-8")

    gate = validate_release_metadata(tmp_path)

    assert gate.valid is False
    assert {issue.code for issue in gate.issues} == {"missing_release_metadata"}


def test_release_manifest_matches_git_index_canonical_blobs_for_archive_portability():
    blobs = _git_index_release_blobs()
    lines = (RELEASE_ROOT / "manifest-sha256.txt").read_text(encoding="utf-8").splitlines()
    entries = dict(line.split("  ", 1)[::-1] for line in lines)
    expected_paths = set(blobs) - {"manifest-sha256.txt"}

    assert set(entries) == expected_paths
    for relative in sorted(expected_paths):
        assert entries[relative] == hashlib.sha256(blobs[relative]).hexdigest(), relative

    critical_route_metadata = {
        f"{route}/{name}"
        for route in ("02-atomicbot-turboquant", "03-animehacker-tq3-0")
        for name in ("data/route.json", "evidence/manifest-sha256.txt")
    }
    assert critical_route_metadata <= entries.keys()

    tree = subprocess.run(
        ["git", "write-tree"],
        cwd=REPOSITORY_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    archive = subprocess.run(
        ["git", "archive", "--format=tar", tree],
        cwd=REPOSITORY_ROOT,
        check=True,
        capture_output=True,
    ).stdout
    prefix = RELEASE_ROOT.relative_to(REPOSITORY_ROOT).as_posix() + "/"
    with tarfile.open(fileobj=io.BytesIO(archive), mode="r:") as bundle:
        archived = {
            member.name.removeprefix(prefix): bundle.extractfile(member).read()
            for member in bundle.getmembers()
            if member.isfile() and member.name.startswith(prefix)
        }
    assert set(archived) == set(blobs)
    for relative, payload in blobs.items():
        assert archived[relative] == payload, relative


def test_every_published_evidence_source_is_exact_in_raw_git_archive(tmp_path: Path):
    expected = _published_evidence_sources()
    tree = subprocess.run(
        ["git", "write-tree"],
        cwd=REPOSITORY_ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    archive_path = tmp_path / "release.tar"
    subprocess.run(
        ["git", "archive", "--format=tar", f"--output={archive_path}", tree],
        cwd=REPOSITORY_ROOT,
        check=True,
    )
    archived: dict[str, tuple[str, int]] = {}
    with tarfile.open(archive_path, mode="r:") as bundle:
        for member in bundle.getmembers():
            if not member.isfile() or member.name not in expected:
                continue
            payload = bundle.extractfile(member).read()
            archived[member.name] = (hashlib.sha256(payload).hexdigest(), len(payload))

    missing = sorted(set(expected) - set(archived))
    mismatched = sorted(
        relative
        for relative, (digest, size) in expected.items()
        if relative in archived and archived[relative] != (digest, size)
    )

    assert missing == []
    assert mismatched == []


@pytest.mark.parametrize(
    ("mutation", "expected_code"),
    [
        (lambda payload: payload.update(valid=False), "release_readiness_invalid"),
        (lambda payload: payload.update(status="failed"), "release_readiness_invalid"),
        (
            lambda payload: payload["basis"].update(
                raw_git_archive_self_contained=False,
                external_hydration_required=True,
            ),
            "release_readiness_invalid",
        ),
        (
            lambda payload: next(iter(payload["qa"]["pdf_reports"].values())).update(
                page_count=0,
                inspected_pages=[],
            ),
            "pdf_qa_invalid",
        ),
        (
            lambda payload: next(
                iter(payload["qa"]["openvino_workbooks"].values())
            ).update(result="failed"),
            "workbook_qa_invalid",
        ),
        (
            lambda payload: next(
                iter(payload["qa"]["openvino_workbooks"].values())
            ).update(formula_count=0),
            "workbook_qa_invalid",
        ),
        (
            lambda payload: next(
                iter(payload["qa"]["openvino_workbooks"].values())
            ).update(missing_sheet_reference_count=1),
            "workbook_qa_invalid",
        ),
        (
            lambda payload: payload["qa"]["openvino_workbooks"][
                "04-openvino-experimental-fork/reports/"
                "openvino-experimental-fork-results.xlsx"
            ].update(machine_absolute_path_count=1),
            "workbook_qa_invalid",
        ),
    ],
)
def test_release_metadata_gate_fails_closed_on_readiness_mutations(
    canonical_release: Path,
    mutation,
    expected_code: str,
):
    path = canonical_release / "validation/release-readiness.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    mutation(payload)
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8", newline="\n")
    _write_filesystem_manifest(canonical_release)

    gate = validate_release_metadata(canonical_release)

    assert gate.valid is False
    assert expected_code in {issue.code for issue in gate.issues}


def test_release_metadata_gate_rejects_human_summary_disagreement(
    canonical_release: Path,
):
    path = canonical_release / "validation/validation-summary.md"
    text = path.read_text(encoding="utf-8")
    assert (
        "Overall result for release package `unified-final-results-2026-09-01-v2`: **Passed"
        in text
    )
    path.write_text(
        text.replace(
            "Overall result for release package `unified-final-results-2026-09-01-v2`: **Passed",
            "Overall result for release package `unified-final-results-2026-09-01-v2`: **Failed",
        ),
        encoding="utf-8",
        newline="\n",
    )
    _write_filesystem_manifest(canonical_release)

    gate = validate_release_metadata(canonical_release)

    assert gate.valid is False
    assert "validation_summary_mismatch" in {issue.code for issue in gate.issues}


def test_release_metadata_gate_rejects_human_summary_count_disagreement(
    canonical_release: Path,
):
    path = canonical_release / "validation/validation-summary.md"
    text = path.read_text(encoding="utf-8")
    assert "| `release_metadata` | Passed | 0 | 6 |" in text
    path.write_text(
        text.replace(
            "| `release_metadata` | Passed | 0 | 6 |",
            "| `release_metadata` | Passed | 0 | 999 |",
        ),
        encoding="utf-8",
        newline="\n",
    )
    _write_filesystem_manifest(canonical_release)

    gate = validate_release_metadata(canonical_release)

    assert gate.valid is False
    assert "validation_summary_mismatch" in {issue.code for issue in gate.issues}


def test_release_metadata_gate_requires_generated_marker_and_activity(
    canonical_release: Path,
):
    path = canonical_release / "ro-crate-metadata.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    expected_pdf = (
        "01-upstream-llama-cpp/reports/upstream-llama-cpp-report.pdf"
    )
    entity = next(item for item in payload["@graph"] if item.get("@id") == expected_pdf)
    entity.pop("additionalType")
    entity.pop("wasGeneratedBy")
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8", newline="\n")
    _write_filesystem_manifest(canonical_release)

    gate = validate_release_metadata(canonical_release)

    codes = {issue.code for issue in gate.issues}
    assert gate.valid is False
    assert "generated_artifact_marker_missing" in codes
    assert "generated_artifact_activity_missing" in codes
