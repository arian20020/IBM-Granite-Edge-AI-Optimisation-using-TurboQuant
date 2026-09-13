import csv
import hashlib
import subprocess
import tempfile
import unittest
import json
from pathlib import Path

from docx import Document

from scripts.testing.tools import audit_official_openvino_docx as audit_wrapper
from scripts.testing.tools.finalize_official_openvino_comparison_workbook import (
    COMPARISON_SECTION_TITLES,
)
from scripts.testing.campaigns.openvino.docx_audit import (
    audit_comparison_docx,
    audit_docx,
    comparison_audit_profile,
)


class OfficialOpenVINODocxAuditTests(unittest.TestCase):
    def make_comparison_fixture(
        self,
        directory: str,
        *,
        matrix_sha256: str = "a" * 64,
        visible_matrix_sha256: str | None = None,
        omitted_heading: str | None = None,
        current_status: str = "Current - pending PR",
        blank_success_cell: bool = False,
        unsupported_winner: bool = False,
        winner_scenario: str | None = None,
        omit_winner: bool = False,
    ):
        directory_path = Path(directory)
        markdown = directory_path / "comparison.md"
        docx = directory_path / "comparison.docx"
        manifest = directory_path / "Controlled-Workbook-Manifest.csv"
        visible_matrix = visible_matrix_sha256 or matrix_sha256
        lines = [
            "# 04 Official OpenVINO Controlled Retest Workbook v1.9",
            "",
            "Controlled retest revision 1.9 (WR-037).",
            "",
            (
                "Adaptive comparison matrix | "
                f"experiments/synthetic/adaptive.json | SHA-256 {matrix_sha256} | "
                "validated cases=5 | declared contexts=25"
            ),
            "",
            "OV-11 OV-12 OV-13 OV-TQ-21 OV-TQ-22",
        ]
        def comparison_table(index: int):
            if winner_scenario is not None and index == 5:
                contexts = ("512", "1024") if winner_scenario == "incomplete" else ("512",)
                return (
                    (
                        "Test ID", "Context", "Cache route", "Activation",
                        "U8 artifact SHA-256", "Runtime evidence",
                    ),
                    tuple(
                        (
                            test_id,
                            context,
                            route,
                            "CPU STANDARD; observed f32/f32 state",
                            "c" * 64,
                            f"runtime/{test_id}-{context}.json#sha256={'d' * 64}",
                        )
                        for context in contexts
                        for test_id, route in (
                            ("OV-12", "U8 STANDARD"),
                            ("OV-TQ-21", "TBQ4"),
                            ("OV-TQ-22", "TBQ3"),
                        )
                    ),
                )
            if winner_scenario is not None and index == 9:
                contexts = ("512", "1024") if winner_scenario == "incomplete" else ("512",)
                rows = []
                for context in contexts:
                    for test_id, score in (
                        ("OV-12", 7.0),
                        ("OV-TQ-21", 9.0),
                        ("OV-TQ-22", 9.0 if winner_scenario == "tie" else 8.0),
                    ):
                        if winner_scenario == "incomplete" and context == "1024" and test_id == "OV-TQ-22":
                            continue
                        scores = tuple(str(score + offset / 10) for offset in range(6))
                        rows.append(
                            (
                                test_id,
                                context,
                                *scores,
                                str(score + 0.25),
                                str(score + 0.25),
                                str(score),
                                str(score + 0.5),
                                f"quality/{test_id}-{context}.json#sha256={'e' * 64}",
                            )
                        )
                return (
                    (
                        "Test ID", "Context", "P1", "P2", "P3", "P4", "P5", "P6",
                        "Mean", "Median", "Minimum", "Maximum", "Quality evidence",
                    ),
                    tuple(rows),
                )
            return (("Identity/context", "Result"), ((f"OV-11/{index}", "Passed"),))

        for index, heading in enumerate(COMPARISON_SECTION_TITLES, start=5):
            headers, rows = comparison_table(index)
            lines.extend(
                (
                    "",
                    f"# {index}. {heading}",
                    "",
                    "| " + " | ".join(headers) + " |",
                    "| " + " | ".join("---" for _ in headers) + " |",
                    *("| " + " | ".join(row) + " |" for row in rows),
                )
            )
        if unsupported_winner:
            lines.extend(("", "Overall winner: **OV-TQ-21**"))
        if winner_scenario is not None and not omit_winner:
            lines.extend(
                ("", "Overall winner: **OV-TQ-21** (shared complete numeric quality).")
            )
        markdown.write_text("\n".join(lines) + "\n", encoding="utf-8")

        document = Document()
        document.add_paragraph(
            "04 Official OpenVINO Controlled Retest Workbook v1.9"
        )
        document.add_paragraph("Controlled retest revision 1.9 (WR-037).")
        document.add_paragraph(
            "Adaptive comparison matrix | experiments/synthetic/adaptive.json | "
            f"SHA-256 {visible_matrix} | validated cases=5 | declared contexts=25"
        )
        document.add_paragraph("OV-11 OV-12 OV-13 OV-TQ-21 OV-TQ-22")
        history = document.add_table(rows=1, cols=7)
        headers = (
            "Version", "Date", "Changed by", "Change", "Affected test IDs",
            "Reference", "Status",
        )
        for cell, value in zip(history.rows[0].cells, headers):
            cell.text = value
        for cell, value in zip(
            history.add_row().cells,
            (
                "1.9", "2026-08-01", "Student and project tooling", "Adaptive comparison",
                "OV-11; OV-12; OV-13; OV-TQ-21; OV-TQ-22; P1-P6", "WR-037",
                current_status,
            ),
        ):
            cell.text = value
        for index, heading in enumerate(COMPARISON_SECTION_TITLES, start=5):
            if heading != omitted_heading:
                document.add_paragraph(f"{index}. {heading}")
            headers, rows = comparison_table(index)
            table = document.add_table(rows=1, cols=len(headers))
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = value
            for row_index, values in enumerate(rows, start=1):
                cells = table.add_row().cells
                for cell, value in zip(cells, values):
                    cell.text = (
                        "" if blank_success_cell and index == 5 and row_index == 1 and value == "Passed"
                        else value
                    )
        if unsupported_winner:
            document.add_paragraph("Overall winner: OV-TQ-21")
        if winner_scenario is not None and not omit_winner:
            document.add_paragraph(
                "Overall winner: OV-TQ-21 (shared complete numeric quality)."
            )
        document.save(docx)
        with manifest.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["Workbook_ID", "Revision", "Last_Validated_DOCX_SHA256"],
            )
            writer.writeheader()
            writer.writerow(
                {
                    "Workbook_ID": "WB-04",
                    "Revision": "1.9",
                    "Last_Validated_DOCX_SHA256": hashlib.sha256(
                        docx.read_bytes()
                    ).hexdigest(),
                }
            )
        return docx, manifest, markdown

    def test_comparison_profile_is_explicit_and_rejects_invalid_matrix_digest(self):
        profile = comparison_audit_profile("a" * 64)
        self.assertEqual(profile.workbook_version, "1.9")
        self.assertEqual(profile.revision_id, "WR-037")
        self.assertEqual(profile.revision_date, "2026-08-01")
        self.assertEqual(profile.current_status, "Current - pending PR")
        self.assertEqual(profile.comparison_headings, COMPARISON_SECTION_TITLES)
        self.assertEqual(
            profile.expected_ids,
            frozenset({"OV-11", "OV-12", "OV-13", "OV-TQ-21", "OV-TQ-22"}),
        )
        for digest in ("a" * 63, "A" * 64, "g" * 64):
            with self.subTest(digest=digest):
                with self.assertRaisesRegex(ValueError, "lowercase SHA-256"):
                    comparison_audit_profile(digest)

    def test_v19_comparison_fixture_passes_profile_and_exact_markdown_table_audit(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, markdown = self.make_comparison_fixture(directory)
            report = audit_comparison_docx(
                docx,
                manifest,
                markdown,
                comparison_audit_profile("a" * 64),
            )
        self.assertTrue(report["accepted"])
        self.assertEqual(report["visible_revision"], "1.9")
        self.assertEqual(report["comparison_heading_count"], 7)
        self.assertEqual(report["comparison_table_count"], 7)

    def test_v19_comparison_audit_rejects_wrong_matrix_heading_blank_status_and_winner(self):
        cases = (
            ({"visible_matrix_sha256": "b" * 64}, "matrix SHA-256"),
            ({"omitted_heading": COMPARISON_SECTION_TITLES[0]}, "heading"),
            ({"blank_success_cell": True}, "blank"),
            ({"current_status": "Current - pending merge"}, "current Status"),
            ({"unsupported_winner": True}, "winner"),
        )
        for arguments, message in cases:
            with self.subTest(arguments=arguments):
                with tempfile.TemporaryDirectory() as directory:
                    docx, manifest, markdown = self.make_comparison_fixture(
                        directory, **arguments
                    )
                    with self.assertRaisesRegex(ValueError, message):
                        audit_comparison_docx(
                            docx,
                            manifest,
                            markdown,
                            comparison_audit_profile("a" * 64),
                        )

    def test_v19_winner_is_independently_recomputed_not_accepted_by_md_docx_parity(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, markdown = self.make_comparison_fixture(
                directory, winner_scenario="eligible"
            )
            report = audit_comparison_docx(
                docx, manifest, markdown, comparison_audit_profile("a" * 64)
            )
            self.assertEqual(report["comparison_winner"], "OV-TQ-21")

        for scenario in ("tie", "incomplete"):
            with self.subTest(scenario=scenario):
                with tempfile.TemporaryDirectory() as directory:
                    docx, manifest, markdown = self.make_comparison_fixture(
                        directory, winner_scenario=scenario
                    )
                    with self.assertRaisesRegex(ValueError, "winner"):
                        audit_comparison_docx(
                            docx,
                            manifest,
                            markdown,
                            comparison_audit_profile("a" * 64),
                        )

    def test_v19_winner_presence_is_derived_and_no_winner_requires_tie_or_incomplete_data(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, markdown = self.make_comparison_fixture(
                directory,
                winner_scenario="eligible",
                omit_winner=True,
            )
            with self.assertRaisesRegex(ValueError, "winner"):
                audit_comparison_docx(
                    docx,
                    manifest,
                    markdown,
                    comparison_audit_profile("a" * 64),
                )

        for scenario in ("tie", "incomplete"):
            with self.subTest(scenario=scenario):
                with tempfile.TemporaryDirectory() as directory:
                    docx, manifest, markdown = self.make_comparison_fixture(
                        directory,
                        winner_scenario=scenario,
                        omit_winner=True,
                    )
                    report = audit_comparison_docx(
                        docx,
                        manifest,
                        markdown,
                        comparison_audit_profile("a" * 64),
                    )
                    self.assertIsNone(report["comparison_winner"])

    def test_wrapper_comparison_release_uses_explicit_markdown_and_matrix_profile(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, markdown = self.make_comparison_fixture(directory)
            output = Path(directory) / "comparison-audit.json"
            result = audit_wrapper.main(
                [
                    "--comparison-release",
                    "--docx", str(docx),
                    "--manifest", str(manifest),
                    "--markdown", str(markdown),
                    "--matrix-sha256", "a" * 64,
                    "--output", str(output),
                ]
            )
            self.assertEqual(result, 0)
            report = json.loads(output.read_text(encoding="utf-8"))
            self.assertTrue(report["accepted"])
            self.assertEqual(report["visible_revision"], "1.9")

    def test_wrapper_comparison_release_rejects_missing_digest_or_legacy_mode_mix(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, markdown = self.make_comparison_fixture(directory)
            common = [
                "--comparison-release",
                "--docx", str(docx),
                "--manifest", str(manifest),
                "--markdown", str(markdown),
                "--output", str(Path(directory) / "audit.json"),
            ]
            self.assertEqual(audit_wrapper.main(common), 1)
            self.assertEqual(
                audit_wrapper.main([*common, "--matrix-sha256", "a" * 64, "--release"]),
                1,
            )

    def make_auditable_fixture(
        self,
        directory: str,
        *,
        presentation_suffix: str = "",
        include_quality_disclosure: bool = True,
        quality_score_cells: tuple[str, str] | None = None,
        include_wr036: bool = True,
        revision_headers: tuple[str, ...] | None = None,
        current_version: str = "1.8",
        current_date: str = "2026-07-31",
        current_status: str = "Current - pending merge",
        duplicate_current: bool = False,
        revision_reference: str = "WR-036",
    ):
        root = Path(__file__).resolve().parents[4]
        matrix = root / "experiments/manifests/official-openvino/retest-matrix.json"
        docx = Path(directory) / "04_Official_OpenVINO_Controlled_Retest_Workbook_v1.docx"
        document = Document()
        document.add_paragraph("04 Official OpenVINO Controlled Retest Workbook v1.8")
        if include_wr036:
            document.add_paragraph("Controlled retest revision 1.8 (WR-036).")
        document.add_paragraph("Document revision history")
        document.add_paragraph(
            " ".join(case.test_id for case in audit_wrapper.load_matrix(matrix))
        )
        presentation_lines = [
            "Accepted formal runtime measurements",
            "OV-TQ-13 | 512",
            "OV-TQ-14 | 512",
            "OV-TQ-14 | 2048",
            "Tests that did not complete",
            "Diagnostic only; no formal benchmark",
            "Missing validated FP16 artifact",
            "RAM safety floor reached",
            "Larger host required",
            "Strict activation proof incomplete",
            "Governed quality campaign stopped at the RAM floor",
        ]
        if include_quality_disclosure:
            presentation_lines.append(
                "No numeric quality score exists and there is no winner."
            )
        presentation_lines.append(presentation_suffix)
        document.add_paragraph("\n".join(presentation_lines))
        expected_headers = (
            "Version", "Date", "Changed by", "Change", "Affected test IDs",
            "Reference", "Status",
        )
        headers = revision_headers or expected_headers
        history = document.add_table(rows=1, cols=len(headers))
        for cell, value in zip(history.rows[0].cells, headers):
            cell.text = value
        for values in (
            ("1.7", "2026-07-31", "Student", "Prior", "OV-*", "WR-035", "Superseded"),
            (
                current_version,
                current_date,
                "Student",
                "Success-focused release",
                "OV-*",
                revision_reference,
                current_status,
            ),
        ):
            for cell, value in zip(history.add_row().cells, values):
                cell.text = value
        if duplicate_current:
            values = (
                "1.8", "2026-07-31", "Student", "Duplicate", "OV-*",
                "WR-036", "Current - pending merge",
            )
            for cell, value in zip(history.add_row().cells, values):
                cell.text = value
        for index in range(1, 10):
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = f"table-{index}"
        if quality_score_cells is not None:
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = quality_score_cells
        document.save(docx)
        manifest = Path(directory) / "Controlled-Workbook-Manifest.csv"
        with manifest.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["Workbook_ID", "Revision", "Last_Validated_DOCX_SHA256"],
            )
            writer.writeheader()
            writer.writerow(
                {
                    "Workbook_ID": "WB-04",
                    "Revision": "1.8",
                    "Last_Validated_DOCX_SHA256": hashlib.sha256(docx.read_bytes()).hexdigest(),
                }
            )
        return docx, manifest, matrix

    def test_v18_fixture_passes_control_and_presentation_audit(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(directory)
            result = audit_docx(
                docx,
                manifest,
                {case.test_id for case in audit_wrapper.load_matrix(matrix)},
            )
        self.assertTrue(result["accepted"])
        self.assertEqual(result["blank_table_cells"], 0)
        self.assertEqual(result["visible_revision"], "1.8")
        self.assertEqual(result["table_count"], 10)
        self.assertEqual(result["presentation_measured_row_count"], 3)
        self.assertEqual(result["presentation_limitation_bullet_count"], 6)
        self.assertEqual(result["revision_history_current_rows"], 1)

    def test_revision_history_requires_visible_wr036_and_exact_table_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory, include_wr036=False, revision_reference="WR-035"
            )
            with self.assertRaisesRegex(ValueError, "WR-036"):
                audit_docx(
                    docx,
                    manifest,
                    {case.test_id for case in audit_wrapper.load_matrix(matrix)},
                )
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory,
                revision_headers=(
                    "Version", "Date", "Changed by", "Summary", "Affected test IDs",
                    "Reference", "Status",
                ),
            )
            with self.assertRaisesRegex(ValueError, "revision-history table"):
                audit_docx(
                    docx,
                    manifest,
                    {case.test_id for case in audit_wrapper.load_matrix(matrix)},
                )

    def test_revision_history_rejects_stale_mismatched_or_duplicate_current_state(self):
        mutations = (
            ({"current_version": "1.7"}, "current Version"),
            ({"current_date": "2026-07-30"}, "current Date"),
            ({"current_status": "Current"}, "current Status"),
            ({"duplicate_current": True}, "exactly one current row"),
        )
        for arguments, message in mutations:
            with self.subTest(arguments=arguments):
                with tempfile.TemporaryDirectory() as directory:
                    docx, manifest, matrix = self.make_auditable_fixture(
                        directory, **arguments
                    )
                    with self.assertRaisesRegex(ValueError, message):
                        audit_docx(
                            docx,
                            manifest,
                            {
                                case.test_id
                                for case in audit_wrapper.load_matrix(matrix)
                            },
                        )

    def test_presentation_rejects_prohibited_claim(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory,
                presentation_suffix="all tests passed",
            )
            with self.assertRaisesRegex(ValueError, "prohibited presentation claim"):
                audit_docx(
                    docx,
                    manifest,
                    {case.test_id for case in audit_wrapper.load_matrix(matrix)},
                )

    def test_presentation_rejects_numeric_quality_score(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory,
                presentation_suffix="Quality score: 8.5 / 10",
            )
            with self.assertRaisesRegex(ValueError, "numeric quality score"):
                audit_docx(
                    docx,
                    manifest,
                    {case.test_id for case in audit_wrapper.load_matrix(matrix)},
                )

    def test_presentation_rejects_pipe_delimited_numeric_quality_score(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory,
                presentation_suffix="Quality score | 8.5 / 10",
            )
            with self.assertRaisesRegex(ValueError, "numeric quality score"):
                audit_docx(
                    docx,
                    manifest,
                    {case.test_id for case in audit_wrapper.load_matrix(matrix)},
                )

    def test_presentation_accepts_disclosure_followed_by_unrelated_version(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory,
                presentation_suffix="Version 1.8",
            )
            result = audit_docx(
                docx,
                manifest,
                {case.test_id for case in audit_wrapper.load_matrix(matrix)},
            )
        self.assertTrue(result["accepted"])

    def test_presentation_rejects_quality_score_split_across_docx_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory,
                quality_score_cells=("Quality score", "8.5 / 10"),
            )
            with self.assertRaisesRegex(ValueError, "numeric quality score"):
                audit_docx(
                    docx,
                    manifest,
                    {case.test_id for case in audit_wrapper.load_matrix(matrix)},
                )

    def test_presentation_rejects_hyphenated_quality_score_across_docx_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory,
                quality_score_cells=("Quality-score:", "8.5 / 10"),
            )
            with self.assertRaisesRegex(ValueError, "numeric quality score"):
                audit_docx(
                    docx,
                    manifest,
                    {case.test_id for case in audit_wrapper.load_matrix(matrix)},
                )

    def test_presentation_requires_the_exact_no_score_no_winner_disclosure(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(
                directory,
                include_quality_disclosure=False,
            )
            with self.assertRaisesRegex(ValueError, "no-score/no-winner disclosure"):
                audit_docx(
                    docx,
                    manifest,
                    {case.test_id for case in audit_wrapper.load_matrix(matrix)},
                )

    def test_blank_table_cell_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad.docx"
            document = Document()
            document.add_paragraph("v1.4 Revision history OV-01")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = ""
            document.save(path)
            with self.assertRaisesRegex(ValueError, "blank DOCX table cells"):
                audit_docx(path, Path("unused.csv"), {"OV-01"})

    def test_wrapper_accepts_custom_paths_and_atomically_writes_report(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(directory)
            output = Path(directory) / "release" / "audit.json"
            result = audit_wrapper.main(
                [
                    "--docx", str(docx),
                    "--manifest", str(manifest),
                    "--matrix", str(matrix),
                    "--output", str(output),
                ]
            )
            self.assertEqual(result, 0)
            report = json.loads(output.read_text(encoding="utf-8"))
            self.assertTrue(report["accepted"])
            self.assertFalse(list(output.parent.glob(".audit.json.*.tmp")))

    def test_wrapper_optional_release_checks_refuse_current_v16_shape(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(directory)
            output = Path(directory) / "audit.json"
            result = audit_wrapper.main(
                [
                    "--docx", str(docx),
                    "--manifest", str(manifest),
                    "--matrix", str(matrix),
                    "--output", str(output),
                    "--expected-visible-revision", "1.7",
                    "--expected-table-count", "10",
                ]
            )
            self.assertEqual(result, 1)
            self.assertFalse(output.exists())

    def test_wrapper_optional_table_count_check_refuses_wrong_count(self):
        with tempfile.TemporaryDirectory() as directory:
            docx, manifest, matrix = self.make_auditable_fixture(directory)
            output = Path(directory) / "audit.json"
            result = audit_wrapper.main(
                [
                    "--docx", str(docx),
                    "--manifest", str(manifest),
                    "--matrix", str(matrix),
                    "--output", str(output),
                    "--expected-visible-revision", "1.8",
                    "--expected-table-count", "9",
                ]
            )
            self.assertEqual(result, 1)
            self.assertFalse(output.exists())

    def test_release_mode_refuses_incompatible_output_override(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "audit.json"
            result = audit_wrapper.main(["--release", "--output", str(output)])
            self.assertEqual(result, 1)
            self.assertFalse(output.exists())

    def test_release_mode_refuses_incompatible_expected_revision_override(self):
        result = audit_wrapper.main(
            ["--release", "--expected-visible-revision", "1.6"]
        )
        self.assertEqual(result, 1)

    def test_release_mode_selects_the_complete_controlled_profile(self):
        configuration = audit_wrapper.resolve_audit_configuration(
            audit_wrapper.parse_args(["--release"])
        )
        self.assertEqual(
            configuration,
            (
                audit_wrapper.DEFAULT_DOCX,
                audit_wrapper.DEFAULT_MANIFEST,
                audit_wrapper.DEFAULT_MATRIX,
                audit_wrapper.RELEASE_OUTPUT,
                "1.8",
                10,
            ),
        )
        self.assertEqual(
            audit_wrapper._sha256(audit_wrapper.DEFAULT_MATRIX),
            audit_wrapper.FROZEN_MATRIX_SHA256,
        )

    def test_release_mode_structural_qa_output_is_not_git_ignored(self):
        repository = Path(__file__).resolve().parents[4]
        ignored = subprocess.run(
            [
                "git",
                "check-ignore",
                "--quiet",
                "--no-index",
                str(audit_wrapper.RELEASE_OUTPUT),
            ],
            cwd=repository,
            check=False,
        )

        self.assertEqual(
            ignored.returncode,
            1,
            f"release QA output is ignored: {audit_wrapper.RELEASE_OUTPUT}",
        )


if __name__ == "__main__":
    unittest.main()
